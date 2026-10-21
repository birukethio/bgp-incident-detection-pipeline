from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "experiments" / "final_error_analysis_v1"
DOCX = OUT / "manuscript_tables_final_error_analysis_v1.docx"


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def set_cell_borders(cell, color: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = tc_borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tc_borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def style_cell(cell, header: bool = False, font_size: int = 8) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_borders(cell)
    if header:
        shade(cell, "1F4E79")
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        for run in p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(font_size)
            run.bold = header
            if header:
                run.font.color.rgb = RGBColor(255, 255, 255)


def remove_style_border(style) -> None:
    style_element = style.element
    p_pr = style_element.find(qn("w:pPr"))
    if p_pr is not None:
        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is not None:
            p_pr.remove(p_bdr)


def remove_paragraph_border(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def fmt(value: str, percent: bool = False, integer: bool = False) -> str:
    if value == "" or value.lower() == "nan":
        return "NA"
    try:
        x = float(value)
        if percent:
            return f"{x * 100:.2f}%"
        if integer:
            return f"{int(round(x))}"
        return f"{x:.3f}"
    except Exception:
        return value


def add_table(doc: Document, title: str, rows: list[dict[str, str]], columns: list[tuple[str, str, bool]], widths: list[float]) -> None:
    h = doc.add_paragraph()
    h.add_run(title).bold = True
    table = doc.add_table(rows=1, cols=len(columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, col in enumerate(columns):
        label = col[1]
        table.rows[0].cells[i].text = label
        table.rows[0].cells[i].width = Inches(widths[i])
        style_cell(table.rows[0].cells[i], header=True, font_size=8)
    for idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, col in enumerate(columns):
            key, _, is_percent = col[:3]
            is_integer = bool(col[3]) if len(col) > 3 else False
            cells[i].text = fmt(row.get(key, ""), percent=is_percent, integer=is_integer)
            cells[i].width = Inches(widths[i])
            if idx % 2 == 1:
                shade(cells[i], "F7FBFF")
            style_cell(cells[i], font_size=7)
    doc.add_paragraph()


def main() -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Inches(11)
    sec.page_height = Inches(8.5)
    sec.top_margin = Inches(0.6)
    sec.bottom_margin = Inches(0.6)
    sec.left_margin = Inches(0.6)
    sec.right_margin = Inches(0.6)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(9)
    doc.styles["Title"].font.name = "Arial"
    doc.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    remove_style_border(doc.styles["Title"])

    title = doc.add_paragraph()
    title.style = doc.styles["Title"]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    remove_paragraph_border(title)
    run = title.add_run("Final Incident Level Error Analysis Tables")
    run.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)

    intro = doc.add_paragraph()
    intro.add_run(
        "These tables summarize frozen-prediction error analysis for the terminal Final Holdout. "
        "They are additive diagnostic materials only; no model was retrained, recalibrated, rerun, or modified."
    )

    summary = read_csv(OUT / "manuscript_table_final_holdout_incident_level_error_summary_v1.csv")
    add_table(
        doc,
        "Table A. Final Holdout Incident Level Error Summary",
        summary,
        [
            ("True Class", "True Class", False),
            ("Predicted Class", "Predicted Class", False),
            ("Count", "Count", False, True),
            ("% of True Class", "% of True Class", True),
            ("% Recovered at Top-2", "% Top-2 Rescue", True),
            ("Mean Top-1 Score", "Mean Top-1 Score", False),
            ("Mean Margin", "Mean Margin", False),
        ],
        [1.1, 1.2, 0.7, 1.15, 1.15, 1.25, 1.0],
    )

    reps = read_csv(OUT / "manuscript_table_representative_final_holdout_misclassifications_v1.csv")
    add_table(
        doc,
        "Table B. Representative Final Holdout Misclassifications",
        reps,
        [
            ("selection_criterion", "Selection Criterion", False),
            ("incident_uid", "Incident", False),
            ("true_class", "True", False),
            ("predicted_class", "Predicted", False),
            ("rank_1_score", "Top-1 Score", False),
            ("rank_2_class", "Top-2 Class", False),
            ("confidence_margin_top1_top2", "Margin", False),
            ("true_class_rank", "True Rank", False, True),
            ("relevant_observations", "Relevant Obs.", False, True),
            ("duration_minutes", "Duration Min.", False, True),
        ],
        [1.7, 1.25, 0.75, 0.85, 0.85, 0.9, 0.75, 0.65, 0.9, 0.9],
    )

    note = doc.add_paragraph()
    note.add_run("Interpretation note: ").bold = True
    note.add_run(
        "Feature and evidence associations in these tables are descriptive only. They should not be interpreted as causal explanations."
    )
    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    main()
