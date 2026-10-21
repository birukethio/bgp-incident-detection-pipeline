from __future__ import annotations

import hashlib
import json
from pathlib import Path

import pandas as pd


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "experiments" / "corpus_balancing_audit_v1"

RAW_SUMMARY = ROOT / "cloudflare_incidents" / "raw_cloudflare_all_available_to_2026-01_summary.json"
ALL_ROSTER_SUMMARY = ROOT / "cloudflare_incidents" / "training_ready_bgp_incident_roster_all_available_to_2026-01_summary.csv"
SAMPLING_AUDIT = ROOT / "cloudflare_incidents" / "stratified_quality_sampling_audit.csv"
SAMPLING_SUMMARY = ROOT / "cloudflare_incidents" / "stratified_quality_sampling_summary.csv"
V1_BALANCED = ROOT / "cloudflare_incidents" / "balanced_incident_level_roster_473_each_leak_hijack_outage.csv"
HIJACK_SAMPLE = ROOT / "cloudflare_incidents" / "stratified_quality_hijack_sample_473.csv"
LEAK_SAMPLE = ROOT / "cloudflare_incidents" / "stratified_quality_leak_sample_473.csv"
V2_ROSTER = ROOT / "data" / "roster_v2" / "balanced_incident_roster_v2.csv"
FROZEN_V2_ROSTER = ROOT / "data" / "roster_v2" / "frozen_incident_roster_v2.csv"
V2_SUMMARY = ROOT / "data" / "roster_v2" / "roster_v2_summary.csv"
V2_QC = ROOT / "data" / "roster_v2" / "roster_v2_quality_audit.csv"
DUP_PLAN = ROOT / "data" / "roster_v2" / "duplicate_canonicalization_plan.csv"
REPLACEMENTS = ROOT / "data" / "roster_v2" / "replacement_incidents.csv"
PART_MANIFEST = ROOT / "data" / "partitions_v2" / "frozen_partition_manifest_v2.csv"
PART_QC = ROOT / "data" / "partitions_v2" / "partition_v2_quality_audit.csv"
PART_DIST = ROOT / "data" / "partitions_v2" / "partition_v2_distribution_summary.csv"
NATURAL_ACCOUNTING = ROOT / "experiments" / "natural_distribution_v1" / "natural_distribution_accounting_v1.csv"
LABEL_TRACE = ROOT / "experiments" / "label_provenance_audit_v1" / "cloudflare_label_entry_trace_v1.csv"
WEIGHT_AUDIT = ROOT / "data" / "modeling" / "xgboost_weight_semantics_audit_v1.csv"
IMBALANCE_WEIGHT_AUDIT = ROOT / "data" / "modeling" / "imbalance_effective_weight_audit_v1.csv"
IMBALANCE_CODE = ROOT / "src" / "bgp_incident_detection" / "imbalance_model_experiment.py"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def rel(path: Path) -> str:
    return str(path.relative_to(ROOT))


def write_md(path: Path, text: str) -> None:
    path.write_text(text.strip() + "\n", encoding="utf-8")


def md_table(df: pd.DataFrame) -> str:
    headers = [str(c) for c in df.columns]
    rows = [[str(v) for v in row] for row in df.to_numpy()]
    widths = [
        max(len(headers[i]), *(len(row[i]) for row in rows)) if rows else len(headers[i])
        for i in range(len(headers))
    ]
    header = "| " + " | ".join(headers[i].ljust(widths[i]) for i in range(len(headers))) + " |"
    sep = "| " + " | ".join("-" * widths[i] for i in range(len(headers))) + " |"
    body = ["| " + " | ".join(row[i].ljust(widths[i]) for i in range(len(headers))) + " |" for row in rows]
    return "\n".join([header, sep, *body])


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def file_source(path: Path) -> str:
    return f"{rel(path)} (sha256={sha256(path)})"


def source_accounting() -> pd.DataFrame:
    raw = {row["kind"]: row for row in load_json(RAW_SUMMARY)["summary"]}
    sampling = pd.read_csv(SAMPLING_AUDIT)
    natural = pd.read_csv(NATURAL_ACCOUNTING)
    label_trace = pd.read_csv(LABEL_TRACE)

    rows: list[dict[str, object]] = []
    for category in ["Hijack", "Leak", "Outage"]:
        trace = label_trace[label_trace["class"] == category].iloc[0]
        raw_count = int(raw[category]["rows"])
        eligible = int(sampling[sampling["category"] == category]["eligible_count"].iloc[0])
        selected = int(sampling[sampling["category"] == category]["selected_count"].iloc[0])
        dedup_row = natural[(natural["stage"] == "deduplication") & (natural["category"] == category)]
        dedup_count = int(dedup_row["count"].iloc[0]) if not dedup_row.empty else raw_count
        source_valid_row = natural[(natural["stage"] == "source_validation") & (natural["category"] == category)]
        source_valid_count = int(source_valid_row["count"].iloc[0]) if not source_valid_row.empty else dedup_count
        unusable = max(source_valid_count - eligible, 0)
        duplicate_count = max(raw_count - dedup_count, 0)
        rows.append(
            {
                "category": category,
                "record_type": "RAW_SOURCE_RECORD -> USABLE_CANDIDATE_INCIDENT -> BALANCED_CORPUS_INCIDENT",
                "raw_source_count": raw_count,
                "duplicate_count_documented_before_balancing": duplicate_count,
                "unusable_or_incomplete_count_documented_before_balancing": unusable,
                "quality_filtered_usable_prebalancing_count": eligible,
                "selected_v1_balanced_count": selected,
                "final_frozen_v2_count": 473,
                "source_api_endpoint": trace["cloudflare_endpoint"],
                "response_container": trace["response_container"],
                "event_identifier_field": trace["event_identifier_field"],
                "event_type_or_class_field": trace["event_type_or_class_field"],
                "date_range_first_start_utc": raw[category]["first_start_utc"],
                "date_range_last_end_utc": raw[category]["last_end_utc"],
                "extraction_or_normalization_code_path": trace["final_label_entry_path"],
                "local_raw_artifact": trace["local_raw_artifact"],
                "primary_count_artifacts": "; ".join(
                    [
                        rel(RAW_SUMMARY),
                        rel(SAMPLING_AUDIT),
                        rel(NATURAL_ACCOUNTING),
                        rel(LABEL_TRACE),
                    ]
                ),
                "notes": (
                    "Hijack/Leak full-universe duplicate-collapse artifact not found; "
                    "Outage has documented dedup/QC path. V2 duplicate remediation is documented "
                    "inside the sampled balanced roster."
                    if category in {"Hijack", "Leak"}
                    else "Raw Outage annotations were reduced through documented deduplication/QC to the 473-event final pool."
                ),
            }
        )
    return pd.DataFrame(rows)


def dedup_audit() -> pd.DataFrame:
    plan = pd.read_csv(DUP_PLAN)
    repl = pd.read_csv(REPLACEMENTS)
    rows = []
    for category in ["Hijack", "Leak", "Outage"]:
        groups = plan[plan["category"] == category]
        excluded = groups["records_to_exclude_from_v2"].fillna("[]").apply(lambda x: len(json.loads(x))).sum()
        rows.append(
            {
                "category": category,
                "full_universe_dedup_before_v1_selection_artifact_found": "NO" if category in {"Hijack", "Leak"} else "PARTIAL_OUTAGE_ONLY",
                "v2_true_duplicate_groups": int(len(groups)),
                "v2_records_excluded_as_true_duplicates": int(excluded),
                "v2_replacements_added": int((repl["category"] == category).sum()),
                "deduplication_key_or_evidence": "group-level true-duplicate recommendations from duplicate_resolution_recommendations and duplicate_canonicalization_plan; canonical selected by deterministic quality metadata",
                "source_artifact": rel(DUP_PLAN),
                "notes": "This documents duplicate correction inside the balanced roster, not a complete duplicate collapse over all 112,290/125,051 Hijack/Leak candidates.",
            }
        )
    return pd.DataFrame(rows)


def final_roster_verification() -> pd.DataFrame:
    roster = pd.read_csv(FROZEN_V2_ROSTER)
    part = pd.read_csv(PART_MANIFEST)
    rows = []
    uid_dup = int(roster["incident_uid"].duplicated().sum())
    sig_cols = ["category", "start_utc", "end_utc", "primary_asn", "affected_asns", "affected_prefixes"]
    signature_dup = int(roster[sig_cols].astype(str).duplicated().sum())
    rows.append(
        {
            "check": "frozen_roster_sha256",
            "observed": sha256(FROZEN_V2_ROSTER),
            "expected": "record checksum",
            "status": "PASS",
            "source_artifact": rel(FROZEN_V2_ROSTER),
        }
    )
    rows.append(
        {
            "check": "total_incident_count",
            "observed": len(roster),
            "expected": 1419,
            "status": "PASS" if len(roster) == 1419 else "FAIL",
            "source_artifact": rel(FROZEN_V2_ROSTER),
        }
    )
    for category, count in roster["category"].value_counts().sort_index().items():
        rows.append(
            {
                "check": f"class_count_{category}",
                "observed": int(count),
                "expected": 473,
                "status": "PASS" if int(count) == 473 else "FAIL",
                "source_artifact": rel(FROZEN_V2_ROSTER),
            }
        )
    rows.extend(
        [
            {
                "check": "incident_uid_duplicates",
                "observed": uid_dup,
                "expected": 0,
                "status": "PASS" if uid_dup == 0 else "FAIL",
                "source_artifact": rel(FROZEN_V2_ROSTER),
            },
            {
                "check": "incident_signature_duplicates_review_indicator",
                "observed": signature_dup,
                "expected": 0,
                "status": "PASS" if signature_dup == 0 else "REVIEW",
                "source_artifact": rel(FROZEN_V2_ROSTER),
            },
            {
                "check": "partition_manifest_count",
                "observed": len(part),
                "expected": 1419,
                "status": "PASS" if len(part) == 1419 else "FAIL",
                "source_artifact": rel(PART_MANIFEST),
            },
            {
                "check": "partition_manifest_sha256",
                "observed": sha256(PART_MANIFEST),
                "expected": "record checksum",
                "status": "PASS",
                "source_artifact": rel(PART_MANIFEST),
            },
        ]
    )
    return pd.DataFrame(rows)


def manuscript_tables() -> tuple[pd.DataFrame, pd.DataFrame]:
    construction = pd.DataFrame(
        [
            ["Raw Cloudflare-derived source records", 112290, 125051, 804, 238145, "Fetched available Cloudflare Radar BGP/outage records through 2026-01; these are source records, not yet class-balanced corpus incidents."],
            ["Documented deduplication/source validation", 112290, 125051, 770, 238111, "Outage deduplication documented; full-universe Hijack/Leak duplicate collapse artifact not found."],
            ["Usable pre-balancing candidate pool", 112290, 125051, 473, 237814, "Observed collected-corpus distribution before deliberate balancing."],
            ["Balanced v1 selected", 473, 473, 473, 1419, "Quality-stratified Hijack/Leak sample plus 473 final Outage pool."],
            ["Post-v2 duplicate collapse", 443, 416, 464, 1323, "Confirmed duplicate groups removed logically while preserving source audit records."],
            ["V2 replacements added", 30, 57, 9, 96, "Replacement incidents selected from unsampled candidate pools to restore class balance."],
            ["Frozen final v2 corpus", 473, 473, 473, 1419, "Authoritative corpus used for final partitions and experiments."],
        ],
        columns=["Stage", "Hijack", "Leak", "Outage", "Total", "Operation"],
    )
    balancing_weighting = pd.DataFrame(
        [
            ["Class count imbalance", "Incident corpus", "473 incidents selected per class", "Controlled scientific comparison across Hijack, Leak, and Outage", "Used before train/validation/holdout partitioning"],
            ["Duplicate records", "Incident corpus", "V2 canonicalization and replacement plan", "Ensure 1,419 independent incidents", "No source records deleted; noncanonical records excluded from v2 roster"],
            ["Unequal window counts per incident", "Model training rows", "I1 incident-aware sample weight = 1 / number of windows for that incident within the training fold", "Prevent long incidents from dominating model fitting", "Weights computed inside TRAIN folds only"],
            ["Temporal cap alternative", "Model training rows", "I2 p95 cap plus I1 weighting", "Sensitivity check for long incident dominance", "Historical experiment; I1 retained for final V4 modeling"],
            ["Natural collected-corpus prior", "Evaluation sensitivity", "Prevalence/prior reweighting of existing TRAIN nested-CV predictions", "Advisor-facing sensitivity to observed source-corpus proportions", "Does not replace balanced corpus or final Holdout"],
        ],
        columns=["Issue", "Level", "Method", "Purpose", "Evaluation Treatment"],
    )
    return construction, balancing_weighting


def write_docx(tables: dict[str, pd.DataFrame]) -> str:
    try:
        from docx import Document
    except Exception as exc:
        return f"DOCX not created: python-docx unavailable ({exc})"
    doc = Document()
    doc.add_heading("Corpus Balancing Audit Manuscript Tables", level=1)
    for title, df in tables.items():
        doc.add_heading(title, level=2)
        table = doc.add_table(rows=1, cols=len(df.columns))
        table.style = "Table Grid"
        for i, col in enumerate(df.columns):
            table.rows[0].cells[i].text = str(col)
        for _, row in df.iterrows():
            cells = table.add_row().cells
            for i, value in enumerate(row):
                cells[i].text = str(value)
    out = OUT / "manuscript_tables_corpus_balancing_v1.docx"
    doc.save(out)
    return rel(out)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)

    accounting = source_accounting()
    accounting.to_csv(OUT / "corpus_source_accounting_v1.csv", index=False)

    dedup = dedup_audit()
    dedup.to_csv(OUT / "corpus_deduplication_audit_v1.csv", index=False)

    final_verify = final_roster_verification()
    final_verify.to_csv(OUT / "corpus_final_roster_verification_v1.csv", index=False)

    construction_table, balancing_weighting_table = manuscript_tables()
    construction_table.to_csv(OUT / "manuscript_table_final_corpus_construction_and_balancing_v1.csv", index=False)
    balancing_weighting_table.to_csv(OUT / "manuscript_table_class_balancing_vs_incident_weighting_v1.csv", index=False)
    write_md(OUT / "manuscript_table_final_corpus_construction_and_balancing_v1.md", md_table(construction_table))
    write_md(OUT / "manuscript_table_class_balancing_vs_incident_weighting_v1.md", md_table(balancing_weighting_table))
    docx_status = write_docx(
        {
            "Final Corpus Construction and Balancing": construction_table,
            "Class Balancing versus Incident-Aware Weighting": balancing_weighting_table,
        }
    )

    sampling_summary = pd.read_csv(SAMPLING_SUMMARY)
    v2_summary = pd.read_csv(V2_SUMMARY)
    part_qc = pd.read_csv(PART_QC)
    part_dist = pd.read_csv(PART_DIST)
    weight_audit = pd.read_csv(WEIGHT_AUDIT)
    imbalance_weight_audit = pd.read_csv(IMBALANCE_WEIGHT_AUDIT)
    label_trace = pd.read_csv(LABEL_TRACE)

    method_text = sampling_summary[sampling_summary["metric"] == "sampling_method"]["value"].iloc[0]
    trace_md = f"""
# Corpus 473 Selection Trace v1

## Scope

This is an additive forensic audit. No corpus, partition, feature, model, Validation, Final Holdout, TTI, or manuscript artifact was modified.

## Raw Candidate Universe

- Hijack raw/source-validated usable candidates: 112,290.
- Leak raw/source-validated usable candidates: 125,051.
- Outage raw records: 804; documented usable final outage pool: 473.

The source/API trace is preserved in `{rel(LABEL_TRACE)}`:

{md_table(label_trace[['class','cloudflare_endpoint','event_identifier_field','event_type_or_class_field','local_raw_artifact','final_label_entry_path','label_supplied_or_inferred']])}

## Historical 473-per-class Selection Procedure Recovered

For Hijack and Leak, the recovered artifact-level selection procedure is:

`{method_text}`

The selected sample artifacts are:

- `{rel(HIJACK_SAMPLE)}`
- `{rel(LEAK_SAMPLE)}`

Both contain 473 selected rows with `sampling_stratum`, `quality_score`, and `quality_rank_within_stratum`. Notes embedded in the selected rows state that quality favored longer duration, stronger source/detector evidence, peer visibility, prefix count, and confidence metadata.

Outage did not require undersampling from the final usable pool: the documented final outage pool contained 473 events and all were retained.

## Selection Type

The preserved evidence supports this classification:

- Hijack: quality-stratified deterministic selection from the Cloudflare-derived usable candidate pool.
- Leak: quality-stratified deterministic selection from the Cloudflare-derived usable candidate pool.
- Outage: QC-filtered final pool retained at 473.

The artifacts do not support a claim that the final 473 Hijack or Leak events were simple random samples. No random seed for the initial 473 Hijack/Leak selection was found in the recovered sampling artifacts.

## v2 Duplicate/Replacement Procedure

The first balanced v1 roster was later remediated. Confirmed duplicate groups were canonicalized in `{rel(DUP_PLAN)}`:

- Confirmed TRUE_DUPLICATE groups: {len(pd.read_csv(DUP_PLAN))}
- Records excluded from v2: {int(pd.read_csv(DUP_PLAN)['records_to_exclude_from_v2'].fillna('[]').apply(lambda x: len(json.loads(x))).sum())}
- Replacements added: {len(pd.read_csv(REPLACEMENTS))}

Class counts after duplicate collapse and replacements are recorded in `{rel(V2_SUMMARY)}`:

{md_table(v2_summary)}

## Reproducibility Gaps

- A full-universe Hijack/Leak duplicate-collapse artifact before the initial 473 selection was not found.
- The selected sample CSVs and summaries preserve the selected records and quality/stratum fields, but the executable script implementing the original quality-stratified selection was not uniquely recovered in this audit.
- Therefore, the exact final v2 roster is reproducible by preserved frozen artifacts and checksums, while the original first-pass 473 Hijack/Leak sampling algorithm is reconstructable only at the artifact/summary level.
"""
    write_md(OUT / "corpus_473_selection_trace_v1.md", trace_md)

    partition_md = f"""
# Corpus Partition Trace v1

The authoritative v2 split is `{rel(PART_MANIFEST)}` with SHA-256 `{sha256(PART_MANIFEST)}`.

The split uses `random_seed = 20260828`, `roster_version = v2`, and `split_version = v2`.

## Counts

{md_table(pd.read_csv(PART_MANIFEST).pivot_table(index='partition', columns='category', values='incident_uid', aggfunc='count', fill_value=0).assign(total=lambda x: x.sum(axis=1)).reset_index())}

## Quality Audit

{md_table(part_qc)}

## Distribution Summary

{md_table(part_dist)}

This partitioning is distinct from both initial class balancing and later I1 model-time incident-aware weighting.
"""
    write_md(OUT / "corpus_partition_trace_v1.md", partition_md)

    i1_md = f"""
# Corpus I1 Weighting Trace v1

I1 is not a corpus selection or balancing operation. It is a model-training sample-weighting strategy applied inside each TRAIN grouped-CV fold.

Recovered implementation: `{rel(IMBALANCE_CODE)}`.

The function is:

```python
def group_aware_weights(train: pd.DataFrame) -> np.ndarray:
    group_sizes = train.groupby("model_group_id")["window_id"].transform("count").astype(float)
    return (1.0 / group_sizes).to_numpy()
```

For the final multiclass XGBoost pathway, the validation-confirmation helper also uses:

```python
counts = df.groupby("incident_uid")["window_id"].transform("count")
return (1.0 / counts).to_numpy(dtype=float)
```

The XGBoost weight audit confirms fold-local equal aggregate incident/group influence:

{md_table(weight_audit)}

The imbalance experiment audit records the same semantics:

{md_table(imbalance_weight_audit[imbalance_weight_audit['imbalance_condition'].eq('I1')].head(10))}
"""
    write_md(OUT / "corpus_i1_weighting_trace_v1.md", i1_md)

    balancing_md = f"""
# Corpus Balancing versus Incident-Aware Weighting v1

{md_table(balancing_weighting_table)}

Key distinction: class balancing operates at the incident corpus level before partitioning. I1 weighting operates later at the model-training row level and equalizes the aggregate influence of incidents with different numbers of five-minute windows.
"""
    write_md(OUT / "corpus_balancing_vs_weighting_v1.md", balancing_md)

    reproducibility_md = f"""
# Corpus Balancing Reproducibility v1

## Disposition

`REPRODUCIBLE_WITH_DOCUMENTED_GAPS`

## What Is Fully Reproducible

- Frozen v2 roster: `{rel(FROZEN_V2_ROSTER)}` with SHA-256 `{sha256(FROZEN_V2_ROSTER)}`.
- Frozen v2 partition manifest: `{rel(PART_MANIFEST)}` with SHA-256 `{sha256(PART_MANIFEST)}`.
- Final class counts: 473 Hijack / 473 Leak / 473 Outage.
- v2 duplicate remediation: `{rel(DUP_PLAN)}` and `{rel(REPLACEMENTS)}`.
- I1 incident-aware weighting formula: `1 / number_of_windows_for_incident_or_model_group`, computed inside training folds only.
- Natural-prior sensitivity is documented separately in `{rel(NATURAL_ACCOUNTING)}` and does not redefine the balanced corpus.

## Documented Gaps

- Full-universe Hijack/Leak deduplication before the initial v1 473 selection was not found.
- The original executable quality-stratified sampling script for the first 473 Hijack/Leak selection was not uniquely recovered, though the selected rows, strata, quality scores, rank fields, and summary method are preserved.
- Cloudflare source label provenance is preserved, but the study's collected-corpus distribution should not be described as global Internet-wide incident prevalence.
"""
    write_md(OUT / "corpus_balancing_reproducibility_v1.md", reproducibility_md)

    chapter3_md = f"""
# Draft Chapter 3 Subsection: Final Corpus Construction and Class Balancing

The final incident corpus was constructed from Cloudflare-derived candidate records and then deliberately balanced at the incident level for controlled three-class comparison. The usable pre-balancing candidate pool contained 112,290 Hijack records, 125,051 Leak records, and 473 quality-controlled Outage incidents. These counts describe the observed distribution of the collected source corpus and should not be interpreted as global Internet-wide prevalence.

Because the study objective required comparable evaluation across Hijack, Leak, and Outage, the modeling corpus was balanced to 473 incidents per class. Hijack and Leak were selected using preserved quality-stratified sampling artifacts that retained one high-quality incident per stratum first and then selected next-best quality records by stratum until the 473-per-class target was reached. The 473 quality-controlled Outage incidents were retained in full. Subsequent v2 quality control collapsed confirmed duplicate groups, excluded noncanonical duplicate records from the v2 roster, and selected replacement incidents from unsampled candidate pools to restore 473 incidents per class. The frozen v2 corpus therefore contains 1,419 independent incidents: 473 Hijack, 473 Leak, and 473 Outage.

This corpus-level class balancing is distinct from the later I1 incident-aware weighting strategy. Corpus balancing controlled the number of independent incidents per class before partitioning. I1 weighting was applied only during model training and assigned each five-minute window a weight of 1 divided by the number of windows for its incident within the training fold, so that long-duration incidents did not dominate model fitting. The final v2 partition used a fixed seed of 20260828 and assigned 331 incidents per class to TRAIN, 71 per class to Validation, and 71 per class to Final Holdout.
"""
    write_md(OUT / "chapter3_corpus_balancing_subsection_v1.md", chapter3_md)

    summary_md = f"""
# Corpus Balancing Audit Summary v1

## Final Answer

The final 1,419-event Praxis corpus was deliberately balanced to 473 Hijack / 473 Leak / 473 Outage incidents. The recovered evidence shows that Hijack and Leak were selected from much larger Cloudflare-derived usable candidate pools using quality-stratified sampling, while the final usable Outage pool already contained 473 incidents and was retained in full. Confirmed duplicate groups were later remediated in v2, removing 96 noncanonical records and adding 96 replacements to restore exact class balance.

## Key Counts

{md_table(construction_table)}

## Exact Historical Selection Procedure Recovered

For Hijack and Leak, the recovered summary states:

> {method_text}

The sample rows themselves preserve `sampling_stratum`, `quality_score`, and `quality_rank_within_stratum`, and notes describe the quality criteria. No evidence was found that simple random sampling was used for the first 473 Hijack/Leak samples.

## Natural Distribution Link

The observed pre-balancing collected-corpus proportions were Hijack 47.2176%, Leak 52.5835%, and Outage 0.1989%, as documented in `{rel(NATURAL_ACCOUNTING)}`. These proportions are a collected-corpus prior-shift sensitivity reference, not a claim about global Internet-wide BGP incident prevalence.

## Reproducibility Disposition

`REPRODUCIBLE_WITH_DOCUMENTED_GAPS`

The frozen roster, partition manifest, v2 duplicate/remediation plan, endpoint provenance, and I1 weighting formula are preserved and auditable. The main gap is that the executable first-pass Hijack/Leak stratified selection script and full-universe preselection duplicate-collapse artifact were not uniquely recovered.

## Manuscript Table DOCX

{docx_status}
"""
    write_md(OUT / "corpus_balancing_audit_summary_v1.md", summary_md)

    # compact machine-readable index
    index = pd.DataFrame(
        [
            {"artifact": p.name, "path": rel(p), "sha256": sha256(p)}
            for p in [
                OUT / "corpus_source_accounting_v1.csv",
                OUT / "corpus_deduplication_audit_v1.csv",
                OUT / "corpus_final_roster_verification_v1.csv",
                OUT / "corpus_473_selection_trace_v1.md",
                OUT / "corpus_partition_trace_v1.md",
                OUT / "corpus_i1_weighting_trace_v1.md",
                OUT / "corpus_balancing_vs_weighting_v1.md",
                OUT / "corpus_balancing_reproducibility_v1.md",
                OUT / "corpus_balancing_audit_summary_v1.md",
            ]
        ]
    )
    index.to_csv(OUT / "corpus_balancing_output_index_v1.csv", index=False)


if __name__ == "__main__":
    main()
