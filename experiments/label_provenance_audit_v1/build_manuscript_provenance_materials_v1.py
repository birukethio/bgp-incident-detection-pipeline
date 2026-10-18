#!/usr/bin/env python3
"""Build manuscript-ready label provenance materials.

This script is additive and writes only new v1 manuscript/defense artifacts
inside experiments/label_provenance_audit_v1.
"""

from __future__ import annotations

import csv
import textwrap
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import landscape, letter
from reportlab.pdfgen import canvas
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


BASE = Path(__file__).resolve().parent


def write_csv(path: Path, rows: list[dict[str, str]], fieldnames: list[str]) -> None:
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def write_md_table(path: Path, title: str, caption: str, rows: list[dict[str, str]], fields: list[str]) -> None:
    labels = [field.replace("_", " ").title() for field in fields]
    lines = [f"# {title}", "", caption, "", "|" + "|".join(labels) + "|", "|" + "|".join(["---"] * len(fields)) + "|"]
    for row in rows:
        vals = [str(row.get(field, "")).replace("\n", "<br>") for field in fields]
        lines.append("|" + "|".join(vals) + "|")
    lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, color: str = "000000") -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(text) < 28 else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_borders(cell) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D9D9D9")


def build_docx_table(rows: list[dict[str, str]], fields: list[str]) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)

    title = doc.add_paragraph()
    title.style = doc.styles["Title"]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Incident Label Provenance and Leakage Assessment by Class")
    run.font.color.rgb = RGBColor(0, 0, 0)

    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.LEFT
    intro_run = intro.add_run(
        "This table summarizes how Cloudflare Radar labels entered the final corpus and how the audit "
        "distinguished external label provenance from documented evidence-source overlap."
    )
    intro_run.font.size = Pt(10.5)

    table = doc.add_table(rows=1, cols=len(fields))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [1.0, 1.35, 1.45, 1.45, 2.1, 1.2]

    header = table.rows[0].cells
    for i, field in enumerate(fields):
        set_cell_text(header[i], field.replace("_", " ").title(), bold=True, color="FFFFFF")
        set_cell_shading(header[i], "1F4E79")
        set_cell_borders(header[i])
        header[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        header[i].width = Inches(widths[i])

    for row in rows:
        cells = table.add_row().cells
        for i, field in enumerate(fields):
            set_cell_text(cells[i], row[field])
            if len(table.rows) % 2 == 0:
                set_cell_shading(cells[i], "F7FBFF")
            set_cell_borders(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[i].width = Inches(widths[i])

    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.LEFT
    c_run = caption.add_run(
        "Note. Direct target leakage was not detected. Evidence-source overlap is documented because "
        "Cloudflare and the model both observe manifestations of the underlying routing event."
    )
    c_run.italic = True
    c_run.font.size = Pt(9)

    doc.save(BASE / "manuscript_table_label_provenance_v1.docx")


FONT_DIR = Path(
    "/Users/biruk/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/"
    "libreoffice-headless/libreoffice/LibreOfficeDev.app/Contents/Resources/fonts/truetype"
)


def font(name: str, size: int) -> ImageFont.FreeTypeFont:
    try:
        return ImageFont.truetype(str(FONT_DIR / name), size=size)
    except Exception:
        return ImageFont.load_default()


def wrapped_lines(text: str, draw: ImageDraw.ImageDraw, fnt, max_width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        trial = word if not current else f"{current} {word}"
        if draw.textbbox((0, 0), trial, font=fnt)[2] <= max_width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def draw_centered_text(draw, box_xy, text: str, fnt, fill="#111111", max_width_pad=26):
    x0, y0, x1, y1 = box_xy
    lines = wrapped_lines(text, draw, fnt, (x1 - x0) - max_width_pad)
    line_height = int(fnt.size * 1.22)
    total_h = line_height * len(lines)
    y = y0 + ((y1 - y0) - total_h) / 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=fnt)
        x = x0 + ((x1 - x0) - (bbox[2] - bbox[0])) / 2
        draw.text((x, y), line, font=fnt, fill=fill)
        y += line_height


def draw_box(draw, xy, text, fill, outline="#444444", fnt=None, radius=18, width=3):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)
    draw_centered_text(draw, xy, text, fnt or font("DejaVuSans.ttf", 30))


def draw_arrow(draw, start, end, fill="#555555", width=4):
    draw.line([start, end], fill=fill, width=width)
    import math

    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 18
    spread = 0.45
    p1 = (
        end[0] - length * math.cos(angle - spread),
        end[1] - length * math.sin(angle - spread),
    )
    p2 = (
        end[0] - length * math.cos(angle + spread),
        end[1] - length * math.sin(angle + spread),
    )
    draw.polygon([end, p1, p2], fill=fill)


def build_lineage_figure() -> None:
    scale = 3
    w, h = 2400, 1350
    img = Image.new("RGB", (w, h), "white")
    draw = ImageDraw.Draw(img)
    title_f = font("DejaVuSans-Bold.ttf", 48)
    sub_f = font("DejaVuSans.ttf", 25)
    head_f = font("DejaVuSans-Bold.ttf", 30)
    box_f = font("DejaVuSans.ttf", 24)
    small_f = font("DejaVuSans.ttf", 21)

    def tcenter(y, text, fnt, fill="#111111"):
        bbox = draw.textbbox((0, 0), text, font=fnt)
        draw.text(((w - (bbox[2] - bbox[0])) / 2, y), text, font=fnt, fill=fill)

    tcenter(55, "Label Provenance and Feature Lineage Audit", title_f)
    tcenter(
        125,
        "External class labels and predictive features were constructed through separate branches; documented evidence-source overlap was expected and audited.",
        sub_f,
        "#333333",
    )
    draw.text((550, 230), "External Label Provenance", font=head_f, fill="#111111", anchor="mm")
    draw.text((1795, 230), "V4 Predictive Feature Representation", font=head_f, fill="#111111", anchor="mm")

    draw_box(draw, (170, 345, 930, 455), "Cloudflare Radar", "#DCEAF7", fnt=font("DejaVuSans-Bold.ttf", 29))
    draw_box(draw, (95, 560, 390, 680), "BGP Hijack Events", "#E8F2FB", fnt=box_f)
    draw_box(draw, (425, 560, 720, 680), "BGP Route Leak Events", "#E8F2FB", fnt=box_f)
    draw_box(draw, (755, 560, 1050, 680), "Outage Annotations eventType = OUTAGE", "#E8F2FB", fnt=small_f)
    draw_box(draw, (255, 805, 845, 930), "Externally sourced class label", "#F2F2F2", fnt=box_f)
    draw_box(draw, (310, 1040, 790, 1155), "Final labels Hijack / Leak / Outage", "#E6E6E6", fnt=box_f)

    draw_box(draw, (1365, 345, 1635, 455), "RouteViews BGP Updates", "#E9F5E9", fnt=small_f)
    draw_box(draw, (1685, 345, 1925, 455), "CAIDA AS Relationships", "#E9F5E9", fnt=small_f)
    draw_box(draw, (1970, 345, 2225, 455), "Historical RPKI / ROV", "#E9F5E9", fnt=small_f)
    draw_box(draw, (1460, 570, 2130, 700), "74-feature V4 representation", "#F2FAF2", fnt=box_f)
    draw_box(draw, (1460, 810, 2130, 935), "Random Forest / XGBoost development", "#F2F2F2", fnt=box_f)
    draw_box(draw, (1460, 1040, 2130, 1155), "Predicted class probabilities", "#E6E6E6", fnt=box_f)

    draw_arrow(draw, (550, 455), (245, 560))
    draw_arrow(draw, (550, 455), (572, 560))
    draw_arrow(draw, (550, 455), (902, 560))
    draw_arrow(draw, (245, 680), (430, 805))
    draw_arrow(draw, (572, 680), (550, 805))
    draw_arrow(draw, (902, 680), (670, 805))
    draw_arrow(draw, (550, 930), (550, 1040))

    draw_arrow(draw, (1500, 455), (1600, 570))
    draw_arrow(draw, (1805, 455), (1800, 570))
    draw_arrow(draw, (2100, 455), (1990, 570))
    draw_arrow(draw, (1795, 700), (1795, 810))
    draw_arrow(draw, (1795, 935), (1795, 1040))

    for y1 in range(295, 1025, 45):
        draw.line([(1200, y1), (1200, y1 + 22)], fill="#888888", width=3)
    tcenter(1175, "Audit boundary: label metadata is not part of the 74-feature input", font("DejaVuSans.ttf", 23), "#333333")

    draw_box(
        draw,
        (620, 1210, 1780, 1300),
        "Cloudflare event label / source classification is not included as a predictive feature.",
        "#FFF7D6",
        outline="#8C6D1F",
        fnt=font("DejaVuSans.ttf", 23),
        width=3,
    )
    tcenter(
        1308,
        "Expected evidence-source overlap exists because both branches describe the same underlying routing event; the audit found no direct target leakage.",
        font("DejaVuSans.ttf", 20),
        "#333333",
    )

    png = BASE / "manuscript_figure_label_feature_lineage_v1.png"
    jpg = BASE / "manuscript_figure_label_feature_lineage_v1_highres.jpg"
    pdf = BASE / "manuscript_figure_label_feature_lineage_v1.pdf"
    img.save(png, dpi=(300, 300))
    img.resize((w * scale // 2, h * scale // 2), Image.Resampling.LANCZOS).save(jpg, quality=96, dpi=(450, 450))

    c = canvas.Canvas(str(pdf), pagesize=landscape(letter))
    page_w, page_h = landscape(letter)
    c.drawImage(str(png), 18, 18, width=page_w - 36, height=page_h - 36, preserveAspectRatio=True, anchor="c")
    c.showPage()
    c.save()


def main() -> None:
    provenance_rows = [
        {
            "incident_class": "Hijack",
            "external_label_source": "Cloudflare Radar BGP Hijack Events feed",
            "cloudflare_feed_endpoint": "/radar/bgp/hijacks/events",
            "local_label_assignment": "Assigned Hijack because event came from the hijack feed",
            "evidence_source_overlap": "Documented overlap with public BGP, RPKI/IRR, AS-path, and AS-relationship evidence",
            "direct_target_leakage": "None detected",
        },
        {
            "incident_class": "Leak",
            "external_label_source": "Cloudflare Radar BGP Route Leak Events feed",
            "cloudflare_feed_endpoint": "/radar/bgp/leaks/events",
            "local_label_assignment": "Assigned Leak because event came from the route leak feed",
            "evidence_source_overlap": "Documented overlap with BGP AS-path and AS-relationship/valley-free evidence",
            "direct_target_leakage": "None detected",
        },
        {
            "incident_class": "Outage",
            "external_label_source": "Cloudflare Radar outage annotations",
            "cloudflare_feed_endpoint": "/radar/annotations/outages with eventType = OUTAGE",
            "local_label_assignment": "Assigned Outage because record was an outage annotation and passed local ASN/network quality filters",
            "evidence_source_overlap": "Broader Internet outage annotation; not a BGP-specific outage detector",
            "direct_target_leakage": "None detected",
        },
    ]
    provenance_fields = [
        "incident_class",
        "external_label_source",
        "cloudflare_feed_endpoint",
        "local_label_assignment",
        "evidence_source_overlap",
        "direct_target_leakage",
    ]
    write_csv(BASE / "manuscript_table_label_provenance_v1.csv", provenance_rows, provenance_fields)
    write_md_table(
        BASE / "manuscript_table_label_provenance_v1.md",
        "Incident Label Provenance and Leakage Assessment by Class",
        "Table. This table summarizes how Cloudflare Radar labels entered the final corpus and how the audit separated external label provenance from evidence-source overlap.",
        provenance_rows,
        provenance_fields,
    )
    build_docx_table(provenance_rows, provenance_fields)

    summary_rows = [
        {"audit_item": "Total incidents audited", "count": "1,419", "interpretation": "Final balanced corpus: 473 Hijack, 473 Leak, 473 Outage"},
        {"audit_item": "PARTIAL_EVIDENCE_OVERLAP", "count": "1,419", "interpretation": "Expected conceptual overlap between Cloudflare source evidence and routing/security model features"},
        {"audit_item": "POTENTIAL_CIRCULAR_LABELING", "count": "0", "interpretation": "No case found where the local pipeline inferred the class directly from the same variables supplied as model features"},
        {"audit_item": "DIRECT_TARGET_LEAKAGE", "count": "0", "interpretation": "No class label, source event type, provider, incident ID, or source classification included as a model feature"},
        {"audit_item": "REVIEW_REQUIRED", "count": "0", "interpretation": "No remaining provenance-review cases in the final audit"},
        {"audit_item": "Final disposition", "count": "PASS_WITH_DOCUMENTED_EVIDENCE_SOURCE_OVERLAP", "interpretation": "Labels are externally sourced, with documented evidence-source overlap rather than full evidence independence"},
    ]
    summary_fields = ["audit_item", "count", "interpretation"]
    write_csv(BASE / "manuscript_table_provenance_audit_summary_v1.csv", summary_rows, summary_fields)
    write_md_table(
        BASE / "manuscript_table_provenance_audit_summary_v1.md",
        "Label-Provenance Audit Summary",
        "Table. Compact summary of the final label-provenance and leakage audit.",
        summary_rows,
        summary_fields,
    )

    limitation_rows = [
        {
            "limitation": "Hijack and Leak evidence overlaps conceptually with model evidence",
            "interpretation": "Cloudflare and the model both observe routing behavior associated with the same event, so labels should not be described as fully evidence-independent.",
        },
        {
            "limitation": "Cloudflare uses public BGP and routing-security information",
            "interpretation": "This supports external label provenance but also creates documented overlap with Behavioral, CAIDA, and RPKI/ROV feature families.",
        },
        {
            "limitation": "Externally sourced does not mean fully independent",
            "interpretation": "The labels came from Cloudflare rather than the local model pipeline, but the underlying public routing evidence can be related.",
        },
        {
            "limitation": "Outage labels are Internet outage annotations",
            "interpretation": "The Outage class should be interpreted as ASN/network-usable Cloudflare outage annotations, not as a dedicated BGP-outage detector output.",
        },
        {
            "limitation": "Direct target leakage was not detected",
            "interpretation": "The feature schema did not include label text, source event type, provider identity, source classification, or incident identifiers.",
        },
    ]
    limitation_fields = ["limitation", "interpretation"]
    write_csv(BASE / "manuscript_table_provenance_limitations_v1.csv", limitation_rows, limitation_fields)
    write_md_table(
        BASE / "manuscript_table_provenance_limitations_v1.md",
        "Label-Provenance Limitations and Interpretation",
        "Table. Interpretation limits for using Cloudflare Radar labels as externally sourced ground truth.",
        limitation_rows,
        limitation_fields,
    )

    (BASE / "chapter3_label_provenance_subsection_v1.md").write_text(
        """# Label Provenance and Leakage Control

All incident labels in the final corpus were sourced from Cloudflare Radar before local feature extraction. Hijack labels came from the Cloudflare Radar BGP Hijack Events feed, and Leak labels came from the Cloudflare Radar BGP Route Leak Events feed. Outage labels came from Cloudflare Radar outage annotations with `eventType = OUTAGE`. The outage source is therefore broader than a BGP-specific outage classifier; the local pipeline retained outage annotations with usable timing and ASN/network evidence for RouteViews-based analysis.

The local machine-learning pipeline did not infer Hijack, Leak, or Outage labels from the 74 V4 predictive features. Instead, local code assigned the class because each record originated from the corresponding Cloudflare feed or outage annotation category. The 74-feature V4 representation contained 44 Behavioral features, 14 CAIDA AS-relationship features, and 16 historical RPKI/ROV features; it did not include source classification, event type, source provider, incident ID, label text, or any label-derived numeric encoding.

The audit distinguishes external label provenance from complete evidence-source independence. Cloudflare's published methodologies for BGP hijacks and route leaks use public routing and routing-security evidence, including BGP update data, AS-path or AS-relationship evidence, and RPKI/IRR-style security context. These sources conceptually overlap with some Behavioral, CAIDA, and RPKI/ROV features. The final disposition was `PASS_WITH_DOCUMENTED_EVIDENCE_SOURCE_OVERLAP`: the labels were externally sourced, expected evidence overlap was documented, and no direct target leakage was detected.
""",
        encoding="utf-8",
    )

    (BASE / "chapter5_label_provenance_limitation_v1.md").write_text(
        """# Label Provenance Limitation

The incident labels were externally sourced from Cloudflare Radar, but evidence-source independence is not complete. Cloudflare and this study both rely on public routing and routing-security evidence to characterize BGP behavior, so Hijack and Leak labels have documented conceptual overlap with some model feature families. This overlap is not direct target leakage because no source classification, event type, provider identity, incident identifier, label text, or label-derived encoding was included as a predictive feature. The Outage class has an additional interpretation limit because it comes from Cloudflare Internet outage annotations rather than a BGP-specific outage detector, so the class should be described as ASN/network-usable outage annotations rather than narrow BGP-control-plane outage labels.
""",
        encoding="utf-8",
    )

    (BASE / "defense_slide_label_provenance_v1.md").write_text(
        """# Label Provenance and Leakage Audit

- Labels externally sourced from Cloudflare Radar.
- Hijack / Leak: BGP-specific Cloudflare event feeds.
- Outage: Cloudflare Internet outage annotations (`eventType = OUTAGE`).
- 1,419/1,419 incidents showed expected evidence-source overlap.
- 0 circular labeling cases.
- 0 direct target leakage cases.
- No class label, source event type, provider, incident ID, or source classification included in the 74-feature V4 input.
- Final disposition: `PASS_WITH_DOCUMENTED_EVIDENCE_SOURCE_OVERLAP`.

Presenter note: Evidence overlap is expected because both the external detector and the ML model observe manifestations of the same routing event. The audit distinguishes that conceptual overlap from direct target leakage.
""",
        encoding="utf-8",
    )

    (BASE / "label_provenance_appendix_material_index_v1.md").write_text(
        """# Label Provenance Appendix Material Index

The main chapters should use the compact provenance tables and concise methodology text. The detailed audit artifacts are better suited for an appendix or supporting artifact set.

| Appendix Candidate | Reason | Recommended Placement |
|---|---|---|
| `label_provenance_incident_level_v1.csv` | Full incident-level provenance is too large for the main chapter but supports reproducibility. | Appendix: Label provenance detail |
| `label_provenance_source_summary_v1.csv` | Summarizes source/provider and class relationships. | Appendix: Source summary |
| `label_feature_overlap_v1.csv` | Shows conceptual relationship between label evidence and feature families. | Appendix: Feature-overlap audit |
| `direct_target_leakage_check_v1.csv` | Documents feature-level target-leakage screening. | Appendix: Leakage checklist |
| `cloudflare_label_entry_trace_v1.md` and `.csv` | Traces Cloudflare API endpoints, response fields, and local code path. | Appendix: Cloudflare API lineage |
| `label_provenance_final_disposition_v1.md` | Records the final audit disposition and stop condition. | Appendix: Final provenance disposition |
""",
        encoding="utf-8",
    )

    build_lineage_figure()


if __name__ == "__main__":
    main()
