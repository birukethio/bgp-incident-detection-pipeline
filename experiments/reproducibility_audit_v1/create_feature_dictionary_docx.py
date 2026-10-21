from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "experiments" / "reproducibility_audit_v1"
CSV_PATH = OUT / "reproducibility_feature_dictionary_v1.csv"
DOCX_PATH = OUT / "reproducibility_feature_dictionary_appendix_v1.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_borders(cell, color: str = "D9D9D9") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def remove_paragraph_border(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def remove_style_border(style) -> None:
    style_element = style.element
    p_pr = style_element.find(qn("w:pPr"))
    if p_pr is not None:
        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is not None:
            p_pr.remove(p_bdr)


def style_cell(cell, bold: bool = False, fill: str | None = None, font_size: int = 8) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if fill:
        set_cell_shading(cell, fill)
    set_cell_borders(cell)
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.05
        for run in paragraph.runs:
            run.font.name = "Arial"
            run.font.size = Pt(font_size)
            run.bold = bold


def main() -> None:
    with CSV_PATH.open("r", encoding="utf-8", newline="") as f:
        rows = list(csv.DictReader(f))

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    remove_style_border(styles["Title"])

    title = doc.add_paragraph()
    title.style = doc.styles["Title"]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    remove_paragraph_border(title)
    run = title.add_run("Appendix Complete Frozen V4 Feature Dictionary")
    run.font.name = "Arial"
    run.font.size = Pt(18)
    run.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)

    intro = doc.add_paragraph()
    intro.add_run(
        "This appendix documents the 74 frozen V4 features used by the final Praxis model. "
        "The table is generated from the preserved behavioral, CAIDA, and RPKI schema artifacts. "
        "It is an audit document only and does not redefine or modify any feature."
    )

    summary = doc.add_paragraph()
    summary.add_run("Feature families: ").bold = True
    summary.add_run("44 Behavioral, 14 CAIDA AS-relationship, and 16 RPKI/ROV features.")

    columns = ["Order", "Feature", "Family", "Definition", "Formula", "Missing or zero behavior"]
    table = doc.add_table(rows=1, cols=len(columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [0.45, 2.05, 1.25, 3.35, 2.15, 2.15]

    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    prevent_row_split(hdr)
    for idx, text in enumerate(columns):
        cell = hdr.cells[idx]
        cell.width = Inches(widths[idx])
        cell.text = text
        style_cell(cell, bold=True, fill="1F4E79", font_size=8)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.color.rgb = None

    for row in rows:
        new_row = table.add_row()
        prevent_row_split(new_row)
        cells = new_row.cells
        values = [
            row["global_order"],
            row["feature_name"],
            row["feature_family"],
            row["definition"],
            row["formula"],
            row["missing_or_zero_behavior"],
        ]
        for idx, value in enumerate(values):
            cells[idx].width = Inches(widths[idx])
            cells[idx].text = value
            style_cell(cells[idx], fill="F7FBFF" if int(row["global_order"]) % 2 == 0 else None, font_size=7)

    note = doc.add_paragraph()
    note.add_run("Source schemas: ").bold = True
    note.add_run(
        "config/behavioral_feature_schema_v1.yaml; config/caida_feature_schema_v1.yaml; "
        "config/rpki_feature_schema_v1.yaml."
    )

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
