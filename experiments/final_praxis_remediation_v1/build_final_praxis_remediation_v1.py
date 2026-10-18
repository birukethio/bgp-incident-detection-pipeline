from __future__ import annotations

import csv
import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd


REPO = Path(__file__).resolve().parents[2]
OUT = REPO / "experiments" / "final_praxis_remediation_v1"
SOURCE_REQUEST = Path("/Users/biruk/.codex/attachments/567d7cae-8d56-4833-ac55-5f15b131051d/pasted-text.txt")


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def rel(path: str | Path) -> str:
    p = Path(path)
    try:
        return str(p.relative_to(REPO))
    except ValueError:
        return str(p)


def read_text(path: str | Path, limit: int | None = None) -> str:
    p = Path(path)
    if not p.exists():
        return ""
    text = p.read_text(encoding="utf-8", errors="ignore")
    return text if limit is None else text[:limit]


def write_csv(path: Path, rows: list[dict[str, object]], columns: list[str] | None = None) -> None:
    df = pd.DataFrame(rows)
    if columns:
        df = df.reindex(columns=columns)
    df.to_csv(path, index=False)


def write_md_table(path: Path, title: str, df: pd.DataFrame, intro: str = "") -> None:
    lines = [f"# {title}", ""]
    if intro:
        lines += [intro, ""]
    if df.empty:
        lines += ["No rows recorded.", ""]
    else:
        lines += [markdown_table(df), ""]
    path.write_text("\n".join(lines), encoding="utf-8")


def markdown_table(df: pd.DataFrame) -> str:
    cols = list(df.columns)
    lines = [
        "| " + " | ".join(str(c) for c in cols) + " |",
        "| " + " | ".join("---" for _ in cols) + " |",
    ]
    for _, row in df.iterrows():
        vals = []
        for col in cols:
            value = "" if pd.isna(row[col]) else str(row[col])
            value = value.replace("\n", " ").replace("|", "\\|")
            vals.append(value)
        lines.append("| " + " | ".join(vals) + " |")
    return "\n".join(lines)


def maybe(path: str) -> str:
    return path if (REPO / path).exists() else f"{path} (not found)"


def metric_from_csv(path: str, metric: str, category: str | None = None) -> str:
    p = REPO / path
    if not p.exists():
        return "NOT_FOUND"
    df = pd.read_csv(p)
    if "metric" not in df.columns:
        return "UNAVAILABLE"
    sub = df[df["metric"].astype(str).str.lower() == metric.lower()]
    if category and "category" in df.columns:
        sub = sub[sub["category"].astype(str).str.lower() == category.lower()]
    if sub.empty:
        return "UNAVAILABLE"
    value_col = "value" if "value" in sub.columns else sub.columns[-1]
    return str(sub.iloc[0][value_col])


def build_inventory() -> list[dict[str, object]]:
    packages = [
        {
            "package_name": "Label provenance audit",
            "path": "experiments/label_provenance_audit_v1",
            "purpose": "Trace final 1,419 labels, source systems, feature-overlap categories, and direct leakage checks.",
            "primary_question_answered": "Were labels externally sourced and do any features directly encode the target label?",
            "key_result": "Direct target leakage check found no feature containing class, source classification, label text, or label-derived encoding; provenance/source-overlap limitations documented.",
            "final_disposition": "PASS_WITH_DOCUMENTED_OVERLAP",
            "manuscript_relevance": "Chapter 3 label provenance subsection; Chapter 5 source-dependence limitation; appendix tables.",
            "defense_relevance": "Core answer to label provenance and leakage concern.",
            "authoritative_artifacts": "label_provenance_final_disposition_v1.md; direct_target_leakage_check_v1.csv; cloudflare_label_entry_trace_v1.md",
        },
        {
            "package_name": "Cloudflare label-entry trace",
            "path": "experiments/label_provenance_audit_v1/cloudflare_label_entry_trace_v1.md",
            "purpose": "Trace Cloudflare API endpoints, event fields, and local conversion paths for Hijack, Leak, and Outage.",
            "primary_question_answered": "How did Cloudflare events become final H/L/O labels?",
            "key_result": "Hijack/Leak labels came from Cloudflare Radar BGP event endpoints; Outage came from Radar outage annotations and must be described as Cloudflare outage annotation provenance rather than assumed BGP root cause.",
            "final_disposition": "ADDRESSED_WITH_DOCUMENTED_LIMITATION",
            "manuscript_relevance": "Chapter 3 corpus construction and label provenance.",
            "defense_relevance": "Prevents overclaiming label-source independence.",
            "authoritative_artifacts": "cloudflare_label_entry_trace_v1.csv; cloudflare_label_entry_trace_v1.md",
        },
        {
            "package_name": "Natural distribution prior-shift sensitivity",
            "path": "experiments/natural_distribution_v1",
            "purpose": "Audit observed source corpus proportions and test TRAIN prediction sensitivity under collected-source class priors.",
            "primary_question_answered": "Does the balanced corpus hide performance sensitivity to observed source prevalence?",
            "key_result": "Observed pre-balancing proportions were Hijack 47.2176%, Leak 52.5835%, Outage 0.1989%; reweighted Macro F1 0.722300 and weighted F1 0.916680.",
            "final_disposition": "PREVALENCE_REWEIGHTED_SENSITIVITY_COMPLETE",
            "manuscript_relevance": "Chapter 4 sensitivity analysis; Chapter 5 limitation on natural prevalence claims.",
            "defense_relevance": "Shows balanced evaluation and observed-prior sensitivity are different questions.",
            "authoritative_artifacts": "natural_distribution_metrics_v1.csv; natural_distribution_reweighting_verification_v1.md",
        },
        {
            "package_name": "Early-horizon and TTI experiment",
            "path": "experiments/tti_early_horizon_v2",
            "purpose": "Evaluate Level-1 retrospective early classification at 30/60/120/300/600/900 second checkpoints.",
            "primary_question_answered": "Can the frozen Level-1 model classify incidents before the full five-minute analytical window?",
            "key_result": "Disposition PASS_WITH_DOCUMENTED_EXCLUSIONS; Macro F1 rose from 0.790817 at 30s to 0.898440 at 300s, with small/inconsistent gains beyond five minutes.",
            "final_disposition": "PASS_WITH_DOCUMENTED_EXCLUSIONS",
            "manuscript_relevance": "Chapter 3 TTI method; Chapter 4 TTI results; Chapter 5 future real-time discussion.",
            "defense_relevance": "Direct answer to Time to Identify concern.",
            "authoritative_artifacts": "early_horizon_final_disposition_v2.md; early_horizon_metrics_v2.csv; early_horizon_final_verification_v2.md",
        },
        {
            "package_name": "Corpus balancing audit",
            "path": "experiments/corpus_balancing_audit_v1",
            "purpose": "Document raw source counts, usable pre-balancing counts, duplicate remediation, 473/class selection, and I1 weighting distinction.",
            "primary_question_answered": "Why was a balanced corpus used and how was it selected?",
            "key_result": "Frozen roster verified at 1,419 incidents, 473 per class, zero incident UID duplicates; raw candidates were much larger for Hijack/Leak than Outage.",
            "final_disposition": "PASS_WITH_DOCUMENTED_SOURCE_DISTRIBUTION_LIMITATION",
            "manuscript_relevance": "Chapter 3 dataset construction; appendix source accounting.",
            "defense_relevance": "Core answer to class balancing concern.",
            "authoritative_artifacts": "corpus_balancing_audit_summary_v1.md; corpus_final_roster_verification_v1.csv; corpus_source_accounting_v1.csv",
        },
        {
            "package_name": "Calibration audit",
            "path": "experiments/calibration_audit_v1",
            "purpose": "Assess frozen prediction confidence, Brier score, log loss, ECE, and terminology boundaries.",
            "primary_question_answered": "Can the output be called calibrated probability?",
            "key_result": "Final Holdout Brier 0.116324, log loss 0.215438, ECE 0.042380; recommendation is to use normalized class scores/probabilities with calibration caveat, not calibrated probabilities.",
            "final_disposition": "CALIBRATION_DIAGNOSTIC_COMPLETE_NO_RECALIBRATION",
            "manuscript_relevance": "Chapter 4 diagnostics; Chapter 5 limitation.",
            "defense_relevance": "Prevents unsupported probability calibration claims.",
            "authoritative_artifacts": "calibration_metrics_v1.csv; calibration_terminology_recommendation_v1.md",
        },
        {
            "package_name": "Hijack 68538 forensic audit",
            "path": "experiments/hijack_68538_audit_v1",
            "purpose": "Trace provenance, RouteViews evidence, features, quality gate, prediction, and exclusion sensitivity for the flagged incident.",
            "primary_question_answered": "Was Hijack:68538 a data-quality inconsistency and did it materially affect final results?",
            "key_result": "Disposition QUALITY_GATE_INCONSISTENCY; RouteViews data were present but no incident-relevant match was retained; post-hoc exclusion sensitivity did not replace official results.",
            "final_disposition": "QUALITY_GATE_INCONSISTENCY",
            "manuscript_relevance": "Chapter 5 limitations and appendix forensic note.",
            "defense_relevance": "Specific answer to identified edge case.",
            "authoritative_artifacts": "hijack_68538_final_disposition_v1.md; hijack_68538_exclusion_sensitivity_v1.csv",
        },
        {
            "package_name": "Partition independence audit",
            "path": "experiments/partition_independence_audit_v1",
            "purpose": "Check UID overlap, duplicates, related events, temporal/prefix/ASN/source context across TRAIN/Validation/Holdout.",
            "primary_question_answered": "Did partitioning leak incidents into evaluation?",
            "key_result": "Exact incident UID overlap 0; confirmed duplicate-event pairs across partitions 0; final disposition PASS_WITH_DOCUMENTED_RELATED_EVENTS.",
            "final_disposition": "PASS_WITH_DOCUMENTED_RELATED_EVENTS",
            "manuscript_relevance": "Chapter 3 leakage controls; appendix audit detail.",
            "defense_relevance": "Supports evaluation integrity.",
            "authoritative_artifacts": "partition_independence_final_disposition_v1.md; partition_independence_summary_v1.csv",
        },
        {
            "package_name": "Reproducibility audit",
            "path": "experiments/reproducibility_audit_v1",
            "purpose": "Document acquisition, windowing, features, weighting, Level-1, Level-2, metrics, and software environment.",
            "primary_question_answered": "Can the final methodology be reproduced procedurally?",
            "key_result": "Final specification and feature dictionary were documented; Level-2 is procedurally specified but original trained binary was not persisted.",
            "final_disposition": "PASS_WITH_LEVEL2_ARTIFACT_LIMITATION",
            "manuscript_relevance": "Chapter 3 reproducibility subsection; Appendix C.",
            "defense_relevance": "Core defense response to reproducibility concern.",
            "authoritative_artifacts": "reproducibility_final_disposition_v1.md; reproducibility_final_experimental_specification_v1.md; reproducibility_artifact_map_v1.csv",
        },
        {
            "package_name": "Final incident-level error analysis",
            "path": "experiments/final_error_analysis_v1",
            "purpose": "Analyze frozen prediction errors, confusion taxonomy, confidence/margins, Top-2 rescue, and class-specific error patterns.",
            "primary_question_answered": "What did the final model get wrong and how useful is ranking?",
            "key_result": "Final Holdout had 17 Top-1 errors; 15/17 were rescued by Top-2; dominant confusion HIJACK_TO_LEAK with 9 incidents.",
            "final_disposition": "FINAL_ERROR_ANALYSIS_COMPLETE_WITH_DESCRIPTIVE_LIMITATIONS",
            "manuscript_relevance": "Chapter 4 error analysis; appendix detailed errors.",
            "defense_relevance": "Supports ranked diagnostic interpretation.",
            "authoritative_artifacts": "final_error_analysis_summary_v1.md; final_error_confusion_taxonomy_v1.csv; final_error_top2_rescue_v1.csv",
        },
        {
            "package_name": "Source-separated IEC construction and Level-1 closure",
            "path": "experiments/independent_validation_v2",
            "purpose": "Build source-native non-Cloudflare IEC roster, verify implementation chain, close Level-2 external validation, and attempt Level-1 IEC gate.",
            "primary_question_answered": "Was source-separated external confirmation executable?",
            "key_result": "Roster locked at 63 events; Level-2 external validation closed due missing trained artifact; Level-1 IEC inconclusive because no verified source-native IEC reconstruction implementation was available.",
            "final_disposition": "LEVEL1_SOURCE_SEPARATED_VALIDATION_INCONCLUSIVE; LEVEL2_EXTERNAL_VALIDATION_CLOSED_WITH_DOCUMENTED_LIMITATION",
            "manuscript_relevance": "Chapter 5 external validation limitation; appendix IEC audit.",
            "defense_relevance": "Transparent closure of external validation concern.",
            "authoritative_artifacts": "level1_iec_execution_v1/IEC_LEVEL1_final_disposition_v1.md; level2_external_validation_closure_v1/level2_external_validation_closure_v1.md",
        },
        {
            "package_name": "Cross-source label validation",
            "path": "experiments/cross_source_label_validation_v1",
            "purpose": "Match existing Cloudflare incidents to same-event independent non-Cloudflare source classifications.",
            "primary_question_answered": "Do Cloudflare labels agree with independent source-native labels where same-event corroboration exists?",
            "key_result": "59/1,419 incidents corroborated; exact agreement 1.0 among accepted matches; corroboration heavily skewed toward Outage and limited for Hijack/Leak.",
            "final_disposition": "SUPPORTED_WITH_SAMPLE_LIMITATION",
            "manuscript_relevance": "Chapter 4/5 label validation; appendix H.",
            "defense_relevance": "Addresses label-source validity without overclaiming coverage.",
            "authoritative_artifacts": "cross_source_agreement_metrics_v1.csv; cross_source_validation_summary_v1.md; hijack_leak_targeted_expansion/*",
        },
        {
            "package_name": "Final Praxis cleanup and reference/page audit",
            "path": "deliverables/praxis_final_cleanup",
            "purpose": "Prior manuscript condensation, page compliance, citation/reference audit, and final rendered Praxis artifacts.",
            "primary_question_answered": "What manuscript areas still need formatting, page, citation, and consistency remediation?",
            "key_result": "A 100-page final cleanup render exists, with source and final reference/citation/page compliance audits retained.",
            "final_disposition": "MANUSCRIPT_REMEDIATION_HISTORY_AVAILABLE",
            "manuscript_relevance": "Current manuscript revision planning.",
            "defense_relevance": "Ensures committee-facing document is consistent.",
            "authoritative_artifacts": "Final_Praxis_Reference_Audit.csv; Final_Praxis_Page_Compliance.csv; Final_Praxis_Consistency_Audit.md",
        },
    ]
    for p in packages:
        full = REPO / str(p["path"])
        p["exists"] = "YES" if full.exists() else "NO"
    return packages


def build_feedback_matrix() -> list[dict[str, object]]:
    src = rel(SOURCE_REQUEST)
    concerns = [
        ("PF01", "Cloudflare label provenance", "Labels may be opaque or source-dependent.", "The prior Praxis did not sufficiently trace how Cloudflare events entered the corpus.", "Completed label provenance and Cloudflare label-entry trace.", "experiments/label_provenance_audit_v1/cloudflare_label_entry_trace_v1.md", "Endpoints, fields, event identifiers, and local conversion paths documented.", "Outage labels remain Cloudflare outage annotations; not all are globally independent BGP root-cause labels.", "ADDRESSED_WITH_DOCUMENTED_LIMITATION", "3", "Dataset provenance", "Label provenance table", "Label-feature lineage figure", "Appendix B", "Dataset/labels", "Cloudflare is the corpus source; independent validation is limited and separately documented."),
        ("PF02", "Label and feature dependence", "Features could overlap conceptually with source detection evidence.", "The manuscript risked treating all label evidence as independent from routing features.", "Classified overlap into predictive association, partial evidence overlap, circular risk, target leakage, or review.", "experiments/label_provenance_audit_v1/label_feature_overlap_v1.csv", "Conceptual overlap documented separately from direct target leakage.", "Underlying evidence-source independence is limited for routing-derived incident sources.", "ADDRESSED_WITH_DOCUMENTED_LIMITATION", "3,5", "Leakage controls and limitations", "Overlap summary", "Label-feature lineage figure", "Appendix B", "Label provenance", "Conceptual overlap is not the same as leakage; direct target leakage was separately checked."),
        ("PF03", "Direct target leakage", "Model features might encode incident class directly.", "Prior narrative needed a direct leakage check.", "Audited all 74 features for class/source-label encodings.", "experiments/label_provenance_audit_v1/direct_target_leakage_check_v1.csv", "No direct class, source event type, label text, or label-derived numeric encoding found among model features.", "Audit is based on preserved schema and artifacts.", "ADDRESSED", "3", "Feature leakage controls", "Direct target leakage table", "", "Appendix B/C", "Leakage", "No direct target field is present in V4."),
        ("PF04", "Class balancing", "Balanced classes may not represent source prevalence.", "The original manuscript needed clearer separation between controlled balanced evaluation and observed source distribution.", "Completed corpus balancing audit and natural-prior sensitivity.", "experiments/corpus_balancing_audit_v1/corpus_balancing_audit_summary_v1.md", "Final corpus verified at 473/473/473; source raw counts were much larger for Hijack/Leak than Outage.", "Balanced corpus is a controlled scientific design, not a prevalence estimate.", "ADDRESSED_WITH_DOCUMENTED_LIMITATION", "3,5", "Corpus construction", "Corpus construction table", "", "Appendix B", "Dataset", "Balanced corpus supports per-class comparison; it does not estimate global prevalence."),
        ("PF05", "Observed natural class distribution", "Committee may ask about natural observed proportions.", "Natural distribution was not originally separated from the balanced evaluation.", "Completed collected-source prior-shift sensitivity.", "experiments/natural_distribution_v1/natural_distribution_reweighting_verification_v1.md", "Observed proportions: Hijack 47.2176%, Leak 52.5835%, Outage 0.1989%; reweighted Macro F1 0.722300, weighted F1 0.916680.", "TRAIN-only prevalence-reweighted sensitivity; not global Internet prevalence.", "ADDRESSED_WITH_DOCUMENTED_LIMITATION", "4,5", "Sensitivity analysis", "Prior-shift table", "", "Appendix D", "Sensitivity", "This is best named prevalence-reweighted sensitivity, not natural global distribution validation."),
        ("PF06", "Incident weighting", "Long incidents could dominate training via more windows.", "Need to show why I1 was used.", "Documented I1 incident-aware weighting and imbalance experiments.", "docs/chapter4_incident_imbalance_results_v1.csv", "I1 retained to equalize incident influence while retaining all windows.", "Window-count imbalance remains a design consideration; no per-incident cap was selected for final pipeline.", "ADDRESSED", "3,4", "Imbalance strategy", "Imbalance results table", "Figure 4 imbalance comparison", "Appendix D", "Imbalance", "I1 solves unequal incident influence better than raw window weighting."),
        ("PF07", "Train/development/holdout split", "Evaluation may have leakage or unclear partition roles.", "The manuscript needed stronger partition chronology.", "Froze v2 partitions and documented access chronology.", "docs/data_access_chronology_v1.csv", "TRAIN 993, Validation 213, Holdout 213; Holdout sealed until Phase L.", "K1 was post-Validation TRAIN development; do not call K1 Validation-confirmed.", "ADDRESSED_WITH_DOCUMENTED_LIMITATION", "3,5", "Experimental design", "Partition summary", "Partition diagram", "Appendix C", "Evaluation integrity", "Validation supported earlier Level-1 candidate; Final Holdout is K1's independent estimate."),
        ("PF08", "Partition independence", "Related/duplicate incidents could cross partitions.", "Existing split needed forensic evidence.", "Completed partition independence audit.", "experiments/partition_independence_audit_v1/partition_independence_final_disposition_v1.md", "Exact UID overlap 0 and confirmed duplicate-event pairs across partitions 0; related events documented.", "Related ASNs/prefixes/context recur naturally in BGP and are not alone leakage.", "ADDRESSED_WITH_DOCUMENTED_LIMITATION", "3", "Leakage controls", "Partition independence verification", "", "Appendix C", "Partition leakage", "No confirmed duplicate-event leakage across partitions."),
        ("PF09", "Independent/source-separated confirmation", "Final model may rely on Cloudflare-labeled data only.", "Need external confirmation boundary.", "Built IEC rosters and closed execution transparently.", "experiments/independent_validation_v2/level1_iec_execution_v1/IEC_LEVEL1_final_disposition_v1.md", "Level-1 source-separated validation is inconclusive; no IEC incidents scored due missing verified reconstruction implementation.", "External validation remains a limitation/future work; do not claim IEC performance.", "INCONCLUSIVE", "5", "External validation limitation", "IEC status table", "IEC architecture figure", "Appendix I", "External validation", "We attempted but did not execute source-separated model validation under the strict reproducibility gate."),
        ("PF10", "H1/H2/H3 confirmation", "Hypotheses may be overclaimed.", "H3 external confirmation and Level-2 validation needed boundary language.", "Mapped H1/H2/H3 to actual evidence and Level-2 closure.", "experiments/independent_validation_v2/level2_external_validation_closure_v1/level2_H3_external_status_v1.md", "H3 external status is H3_EXTERNAL_CONFIRMATION_NOT_EXECUTED_LEVEL2_ARTIFACT_NOT_PERSISTED.", "H3 supported by TRAIN/Final Holdout historical evidence, not source-separated external confirmation.", "ADDRESSED_WITH_DOCUMENTED_LIMITATION", "5", "Hypothesis disposition", "Hypothesis disposition table", "", "Appendix I", "Hypotheses", "State exactly which hypotheses are supported and where limitations remain."),
        ("PF11", "Five-minute analytical horizon", "Why five minutes if live environments see updates in seconds?", "Original claims risked conflating five-minute window with real-time TTI.", "Completed early-horizon/TTI experiment and saturation check.", "experiments/tti_early_horizon_v2/early_horizon_final_verification_v2.md", "Performance improved through five minutes; beyond 300s gains were small/inconsistent.", "Five minutes is a practical tradeoff point, not a mathematically proven optimum.", "ADDRESSED_WITH_DOCUMENTED_LIMITATION", "3,4,5", "TTI methodology/results", "Horizon performance table", "TTI performance figure", "Appendix G", "TTI", "Five-minute is an analytical tradeoff, not live deployment latency."),
        ("PF12", "TTI", "Time-to-identify was not originally measured.", "The final architecture used complete five-minute windows.", "Completed retrospective checkpoint TTI with 30/60/120/300/600/900 horizons.", "experiments/tti_early_horizon_v2/early_horizon_metrics_v2.csv", "Top-1 cumulative by 300s 910/983; Top-2 cumulative by 300s 965/983; 600/900 exploratory.", "TRAIN-only Level-1 retrospective analysis, not Final Holdout or live TTI.", "ADDRESSED_WITH_DOCUMENTED_LIMITATION", "4,5", "Early classification", "Cumulative TTI table", "Cumulative TTI figure", "Appendix G", "TTI", "Checkpoint TTI is retrospective RouteViews-observable time, not wall-clock enterprise alerting time."),
        ("PF13", "Real-time versus retrospective claims", "System might be described as live deployment-ready.", "Earlier language risked overstating operational deployment.", "Added latency and TTI boundary docs and future Kafka/BMP architecture plan.", "docs/professor_tau_and_latency_answers_v1.md", "25.58 ms downstream latency is computational scoring latency, not full detection latency.", "Live real-time incremental pipeline remains future work.", "ADDRESSED_WITH_DOCUMENTED_LIMITATION", "5", "Future work", "", "Future architecture", "", "Operationalization", "Retrospective RouteViews analysis is not live enterprise deployment."),
        ("PF14", "Level-2 temporal causality", "Retrospective temporal stacker may use completed incident sequence.", "Need to avoid prospective/causal language.", "Independent validation V1 causality audit and Level-2 closure produced limitations.", "experiments/independent_validation_audit_v1/temporal_stacker_causality_audit_v1.md", "Level-2 is retrospective incident-level aggregation, not prospective temporal-causal inference.", "Prospective Level-2 requires future work.", "ADDRESSED_WITH_DOCUMENTED_LIMITATION", "3,5", "Model architecture limitations", "Architecture table", "K1 architecture figure", "Appendix I", "Architecture", "The final K1 stacker is retrospective over configured observation windows."),
        ("PF15", "Calibration/probability terminology", "Probabilities may be overinterpreted as calibrated.", "Need calibration evidence and terminology change.", "Completed frozen-prediction calibration audit.", "experiments/calibration_audit_v1/calibration_terminology_recommendation_v1.md", "Use normalized class probabilities/scores with calibration caveat; no recalibration performed.", "Calibration is diagnostic only.", "ADDRESSED", "4,5", "Calibration diagnostics", "Calibration table", "Reliability figure", "Appendix E", "Calibration", "Do not claim calibrated probabilities."),
        ("PF16", "Hijack:68538", "Specific incident had no directly relevant RouteViews observations.", "Could undermine quality gates.", "Completed forensic audit and exclusion sensitivity.", "experiments/hijack_68538_audit_v1/hijack_68538_final_disposition_v1.md", "Disposition QUALITY_GATE_INCONSISTENCY; official result unchanged.", "Document as limitation rather than remove post hoc.", "ADDRESSED_WITH_DOCUMENTED_LIMITATION", "5", "Data quality limitation", "Exclusion sensitivity table", "", "Appendix B/F", "Data quality", "The case is a documented quality-gate caveat, not a reason to alter final holdout."),
        ("PF17", "Reproducibility", "Committee may require exact reproducible method.", "Need artifact map and final specs.", "Completed reproducibility audit and feature dictionary.", "experiments/reproducibility_audit_v1/reproducibility_final_experimental_specification_v1.md", "Data acquisition, features, weighting, models, and metrics documented; Level-2 binary not persisted.", "Procedural reproducibility differs from original trained artifact recoverability.", "ADDRESSED_WITH_DOCUMENTED_LIMITATION", "3,5", "Reproducibility", "Feature dictionary", "", "Appendix C", "Reproducibility", "Level-2 artifact limitation is now documented and bounded."),
        ("PF18", "Feature dictionary", "Feature definitions must be transparent.", "Earlier manuscript could not bury 74 features.", "Created full 74-feature reproducibility dictionary.", "experiments/reproducibility_audit_v1/reproducibility_feature_dictionary_v1.csv", "44 Behavioral + 14 CAIDA + 16 RPKI/ROV definitions available.", "Full inventory should live in appendix to protect page budget.", "ADDRESSED", "3", "Feature engineering", "Feature family table", "", "Appendix A", "Feature transparency", "Main text summarizes; appendix carries full feature dictionary."),
        ("PF19", "Incident-level error analysis", "Final accuracy alone is not explanatory.", "Need error taxonomy and Top-2 utility.", "Completed final incident-level error analysis.", "experiments/final_error_analysis_v1/final_error_analysis_summary_v1.md", "17 Final Holdout Top-1 errors; 15/17 rescued by Top-2; dominant confusion Hijack to Leak.", "Associative error factors only; not causal explanation.", "ADDRESSED", "4,5", "Error analysis", "Error summary table", "Confusion/error figures", "Appendix F", "Errors", "Top-2 ranking value is a defensible final contribution."),
        ("PF20", "Root-cause terminology", "Classifier may not prove true root cause.", "Use of root-cause language could overclaim.", "Claims boundary narrows language to incident class/cause ranking.", "experiments/final_praxis_remediation_v1/FINAL_RESEARCH_CLAIMS_BOUNDARY_v1.md", "Root-cause determination is not claimed; ranked incident-cause classification is claimed.", "True operational root cause requires operator context and prospective telemetry.", "MANUSCRIPT_REMEDIATION_REQUIRED", "1,3,5", "Terminology", "", "", "", "Terminology", "Use incident type/class/cause ranking, not definitive root-cause determination."),
        ("PF21", "Enterprise-network claims", "RouteViews public data may not represent enterprise-local state.", "Enterprise framing needed limitation.", "Limitations/future work now distinguish public RouteViews from enterprise BMP/Kafka telemetry.", "docs/final_limitations_v1.csv", "Enterprise-local deployment remains future work.", "Public collector observability varies.", "MANUSCRIPT_REMEDIATION_REQUIRED", "1,5", "Scope and limitations", "", "Future architecture", "", "Scope", "Keep enterprise language as motivation and future deployment context, not proven operational environment."),
        ("PF22", "Page-count compliance", "Praxis may exceed page requirements.", "Prior full source was too long and overly historical.", "Prepared page reduction plan and appendices.", "experiments/final_praxis_remediation_v1/FINAL_PAGE_REDUCTION_PLAN_v1.md", "Target 80-100 pages for Chapters 1-5 with details moved to appendices.", "Requires manuscript editing after this planning package.", "MANUSCRIPT_REMEDIATION_REQUIRED", "All", "Page planning", "", "", "", "Manuscript", "Move audit detail to appendix; keep main chapters answer-driven."),
        ("PF23", "Obsolete methodology/results", "Earlier RF, 49-feature, Stage 1, and IEC attempts may confuse reader.", "Need stale artifact and table/figure selection.", "Phase M and this package identify current vs historical artifacts.", "docs/stale_artifact_register_v1.csv", "Final model is K1 XGB/XGB; Stage 1 auxiliary/unresolved; RF and early feature counts historical.", "Existing manuscript must be checked for stale claims.", "MANUSCRIPT_REMEDIATION_REQUIRED", "3,4", "Consistency cleanup", "", "", "", "Manuscript", "Do not present old experiments as final architecture."),
        ("PF24", "References/citations", "APA7 and source support may be incomplete.", "Need citation audit and targeted literature plan.", "Created final reference remediation plan from existing audits.", "deliverables/praxis_final_cleanup/Final_Praxis_Reference_Audit.csv", "Citation/reference remediation items identified; bibliography not rewritten in this task.", "Requires final manual/automated reference pass.", "MANUSCRIPT_REMEDIATION_REQUIRED", "2,References", "Literature review", "", "", "", "References", "Reference remediation remains a manuscript task, not an experiment."),
        ("PF25", "Figure/table numbering and manuscript/template integrity", "Final document may contain stale numbering or formatting drift.", "Need inventory and insertion plan.", "Created final table/figure inventory and page-reduction plan.", "experiments/final_praxis_remediation_v1/FINAL_TABLE_FIGURE_INVENTORY_v1.csv", "Main figures/tables selected and appendix candidates identified.", "Requires final manuscript editing without shrinking fonts/margins.", "MANUSCRIPT_REMEDIATION_REQUIRED", "All", "Tables and figures", "Inventory", "Inventory", "Appendix J", "Formatting", "Use fewer main visuals and move audit detail to appendices."),
    ]
    cols = [
        "Feedback_ID", "Professor_Feedback", "Underlying_Risk", "Original_Praxis_Problem",
        "Remediation_Action", "Authoritative_Evidence", "Key_Result", "Limitation",
        "Final_Status", "Praxis_Chapter", "Praxis_Section", "Table", "Figure",
        "Appendix_Material", "Defense_Slide", "Defense_Talking_Point",
    ]
    return [dict(zip(cols, row)) for row in concerns]


def build_results_registry() -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []

    def add(metric, value, population, experiment, source, usage, notes=""):
        rows.append({
            "metric": metric,
            "value": value,
            "population": population,
            "experiment": experiment,
            "source_artifact": source,
            "manuscript_usage": usage,
            "notes": notes,
        })

    add("raw_hijack_count", "112290", "Collected Cloudflare source corpus", "Corpus balancing audit", "experiments/corpus_balancing_audit_v1/corpus_source_accounting_v1.csv", "Chapter 3 source accounting", "Raw source records before deliberate balancing; not global prevalence.")
    add("raw_leak_count", "125051", "Collected Cloudflare source corpus", "Corpus balancing audit", "experiments/corpus_balancing_audit_v1/corpus_source_accounting_v1.csv", "Chapter 3 source accounting", "Raw source records before deliberate balancing; not global prevalence.")
    add("raw_outage_count", "804", "Collected Cloudflare source corpus", "Corpus balancing audit", "experiments/corpus_balancing_audit_v1/corpus_source_accounting_v1.csv", "Chapter 3 source accounting", "Cloudflare outage annotations.")
    add("usable_prebalancing_hijack_count", "112290", "Collected source-validated corpus", "Corpus balancing audit", "experiments/corpus_balancing_audit_v1/corpus_source_accounting_v1.csv", "Appendix source accounting", "Full-universe Hijack duplicate-collapse artifact not found; documented as limitation.")
    add("usable_prebalancing_leak_count", "125051", "Collected source-validated corpus", "Corpus balancing audit", "experiments/corpus_balancing_audit_v1/corpus_source_accounting_v1.csv", "Appendix source accounting", "Full-universe Leak duplicate-collapse artifact not found; documented as limitation.")
    add("usable_prebalancing_outage_count", "473", "Collected source-validated corpus", "Corpus balancing audit", "experiments/corpus_balancing_audit_v1/corpus_source_accounting_v1.csv", "Chapter 3 source accounting", "Outage annotations reduced through dedup/QC.")
    add("balanced_count_per_class", "473/473/473", "Frozen v2 incident corpus", "Roster v2 freeze", "data/roster_v2/frozen_incident_roster_v2.csv", "Chapter 3 dataset", "Final corpus is balanced by design.")
    add("final_total_incidents", "1419", "Frozen v2 incident corpus", "Roster v2 freeze", "experiments/corpus_balancing_audit_v1/corpus_final_roster_verification_v1.csv", "Chapter 3 dataset", "473 per class.")
    add("TRAIN_incidents", "993", "Frozen partition", "Partition v2", "docs/authoritative_numbers_register_v1.csv", "Chapter 3/4", "331 per class.")
    add("Validation_incidents", "213", "Frozen partition", "Phase J", "docs/authoritative_numbers_register_v1.csv", "Chapter 3/4", "71 per class.")
    add("Final_Holdout_incidents", "213", "Frozen partition", "Phase L", "config/final_praxis_result_v1.yaml", "Chapter 4", "71 per class.")
    add("duplicate_collapse_removed", "96 total; Hijack 30; Leak 57; Outage 9", "Roster v2 remediation", "Duplicate canonicalization", "experiments/corpus_balancing_audit_v1/corpus_deduplication_audit_v1.csv", "Chapter 3 appendix", "Duplicate correction inside balanced roster.")
    add("behavioral_feature_count", "44", "Frozen V4 schema", "Feature engineering", "docs/authoritative_numbers_register_v1.csv", "Chapter 3", "")
    add("caida_feature_count", "14", "Frozen V4 schema", "Feature engineering", "docs/authoritative_numbers_register_v1.csv", "Chapter 3", "")
    add("rpki_rov_feature_count", "16", "Frozen V4 schema", "Feature engineering", "docs/authoritative_numbers_register_v1.csv", "Chapter 3", "")
    add("total_v4_feature_count", "74", "Frozen V4 schema", "Feature engineering", "docs/authoritative_numbers_register_v1.csv", "Chapter 3/4", "")
    add("final_holdout_macro_f1", "0.920177", "Final Holdout", "Phase L terminal evaluation", "config/final_praxis_result_v1.yaml", "Chapter 4 headline result", "Official immutable terminal result.")
    add("final_holdout_macro_f1_95_ci", "[0.880117, 0.954402]", "Final Holdout", "Phase L bootstrap", "config/final_praxis_result_v1.yaml", "Chapter 4 headline result", "")
    add("final_holdout_top1", "0.920188", "Final Holdout", "Phase L terminal evaluation", "config/final_praxis_result_v1.yaml", "Chapter 4 headline result", "")
    add("final_holdout_top2", "0.990610", "Final Holdout", "Phase L terminal evaluation", "config/final_praxis_result_v1.yaml", "Chapter 4 ranking result", "211/213 ranked true class in Top-2.")
    add("final_holdout_mrr", "0.958529", "Final Holdout", "Phase L terminal evaluation", "config/final_praxis_result_v1.yaml", "Chapter 4 ranking result", "")
    add("final_holdout_hijack_f1", "0.877698", "Final Holdout", "Phase L terminal evaluation", "config/final_praxis_result_v1.yaml", "Chapter 4 per-class", "")
    add("final_holdout_leak_f1", "0.904110", "Final Holdout", "Phase L terminal evaluation", "config/final_praxis_result_v1.yaml", "Chapter 4 per-class", "")
    add("final_holdout_outage_f1", "0.978723", "Final Holdout", "Phase L terminal evaluation", "config/final_praxis_result_v1.yaml", "Chapter 4 per-class", "")
    add("final_holdout_top1_errors", "17/213", "Final Holdout", "Final error analysis", "experiments/final_error_analysis_v1/final_error_analysis_summary_v1.md", "Chapter 4 error analysis", "")
    add("dominant_confusion", "HIJACK_TO_LEAK with 9 incidents", "Final Holdout errors", "Final error analysis", "experiments/final_error_analysis_v1/final_error_analysis_summary_v1.md", "Chapter 4 error analysis", "")
    add("top2_rescue", "15/17 Top-1 errors rescued by Top-2", "Final Holdout errors", "Final error analysis", "experiments/final_error_analysis_v1/final_error_analysis_summary_v1.md", "Chapter 4 ranking utility", "")
    for _, r in pd.read_csv(REPO / "experiments/tti_early_horizon_v2/early_horizon_metrics_v2.csv").iterrows():
        add(f"horizon_{int(r.horizon_seconds)}s_macro_f1", f"{r.macro_f1:.6f}", "TRAIN TTI evaluable cohort", "Early-horizon TTI v2", "experiments/tti_early_horizon_v2/early_horizon_metrics_v2.csv", "Chapter 4 TTI", "Level-1 only; retrospective/checkpoint-based.")
        add(f"horizon_{int(r.horizon_seconds)}s_top1", f"{r.top1_accuracy:.6f}", "TRAIN TTI evaluable cohort", "Early-horizon TTI v2", "experiments/tti_early_horizon_v2/early_horizon_metrics_v2.csv", "Chapter 4 TTI", "")
        add(f"horizon_{int(r.horizon_seconds)}s_top2", f"{r.top2_accuracy:.6f}", "TRAIN TTI evaluable cohort", "Early-horizon TTI v2", "experiments/tti_early_horizon_v2/early_horizon_metrics_v2.csv", "Chapter 4 TTI", "")
        add(f"horizon_{int(r.horizon_seconds)}s_mrr", f"{r.mrr:.6f}", "TRAIN TTI evaluable cohort", "Early-horizon TTI v2", "experiments/tti_early_horizon_v2/early_horizon_metrics_v2.csv", "Chapter 4 TTI", "")
    add("top1_cumulative_tti_300s", "910/983 = 92.57%", "TRAIN TTI evaluable cohort", "Early-horizon TTI v2", "experiments/tti_early_horizon_v2/early_horizon_tti_cumulative_v2.csv", "Chapter 4 TTI", "")
    add("top2_cumulative_tti_300s", "965/983 = 98.17%", "TRAIN TTI evaluable cohort", "Early-horizon TTI v2", "experiments/tti_early_horizon_v2/early_horizon_tti_cumulative_v2.csv", "Chapter 4 TTI", "")
    add("final_holdout_brier", "0.116324", "Final Holdout", "Calibration audit", "experiments/calibration_audit_v1/calibration_metrics_v1.csv", "Chapter 4 diagnostics", "Diagnostic only; no recalibration.")
    add("final_holdout_log_loss", "0.215438", "Final Holdout", "Calibration audit", "experiments/calibration_audit_v1/calibration_metrics_v1.csv", "Chapter 4 diagnostics", "")
    add("final_holdout_ece_equal_width", "0.042380", "Final Holdout", "Calibration audit", "experiments/calibration_audit_v1/calibration_metrics_v1.csv", "Chapter 4 diagnostics", "")
    add("final_holdout_mean_top1_confidence", "0.892056", "Final Holdout", "Calibration audit", "experiments/calibration_audit_v1/calibration_metrics_v1.csv", "Chapter 4 diagnostics", "")
    add("natural_prior_macro_f1", "0.722300", "TRAIN nested-CV predictions reweighted", "Natural prior sensitivity", "experiments/natural_distribution_v1/natural_distribution_metrics_v1.csv", "Chapter 4 sensitivity", "Prevalence-reweighted sensitivity, not external generalization.")
    add("natural_prior_weighted_f1", "0.916680", "TRAIN nested-CV predictions reweighted", "Natural prior sensitivity", "experiments/natural_distribution_v1/natural_distribution_metrics_v1.csv", "Chapter 4 sensitivity", "")
    add("natural_prior_top1", "0.914426", "TRAIN nested-CV predictions reweighted", "Natural prior sensitivity", "experiments/natural_distribution_v1/natural_distribution_metrics_v1.csv", "Chapter 4 sensitivity", "")
    add("natural_prior_top2", "0.989852", "TRAIN nested-CV predictions reweighted", "Natural prior sensitivity", "experiments/natural_distribution_v1/natural_distribution_metrics_v1.csv", "Chapter 4 sensitivity", "")
    add("natural_prior_mrr", "0.955522", "TRAIN nested-CV predictions reweighted", "Natural prior sensitivity", "experiments/natural_distribution_v1/natural_distribution_metrics_v1.csv", "Chapter 4 sensitivity", "")
    add("cross_source_incidents_searched", "1419", "Frozen v2 corpus", "Cross-source label validation", "experiments/cross_source_label_validation_v1/cross_source_agreement_metrics_v1.csv", "Chapter 4 label validation", "")
    add("cross_source_corroborated", "59", "Matched same-event incidents", "Cross-source label validation", "experiments/cross_source_label_validation_v1/cross_source_agreement_metrics_v1.csv", "Chapter 4 label validation", "0 Hijack, 1 Leak, 58 Outage in original v1.")
    add("cross_source_exact_agreement", "1.0", "Accepted matched incidents", "Cross-source label validation", "experiments/cross_source_label_validation_v1/cross_source_agreement_metrics_v1.csv", "Chapter 4 label validation", "Among accepted unambiguous matches only.")
    add("cross_source_kappa", "1.0", "Accepted matched incidents", "Cross-source label validation", "experiments/cross_source_label_validation_v1/cross_source_agreement_metrics_v1.csv", "Chapter 4 label validation", "Sample coverage limited.")
    add("IEC_source_separated_roster", "63 total; 16 Hijack / 17 Leak / 30 Outage", "IEC V2", "Source-separated IEC construction", "experiments/independent_validation_v2/v2_IEC_final_source_separated_roster.csv", "Chapter 5 external validation limitation", "")
    add("IEC_Level1_external_scoring", "0 evaluable; inconclusive", "IEC V2", "Level-1 IEC execution", "experiments/independent_validation_v2/level1_iec_execution_v1/IEC_LEVEL1_final_disposition_v1.md", "Chapter 5 limitation", "No IEC model scoring was performed.")
    add("IEC_Level2_external_scoring", "not executed", "IEC V2", "Level-2 closure", "experiments/independent_validation_v2/level2_external_validation_closure_v1/level2_external_validation_closure_v1.md", "Chapter 5 limitation", "Original trained Level-2 artifact was not persisted.")
    return rows


def build_claims_boundary() -> str:
    return """# Final Research Claims Boundary v1

This document is the authoritative boundary for the final Praxis revision. Experimental development is closed. No additional modeling, tuning, validation, or result replacement should be performed unless a demonstrable implementation error is discovered.

## Directly Supported Claims

- The final frozen K1 architecture achieved strong incident-level classification and ranking on the untouched Final Holdout: Macro F1 0.920177, Top-1 0.920188, Top-2 0.990610, and MRR 0.958529.
- The final corpus contains 1,419 balanced incidents, 473 each for Hijack, Leak, and Outage, with frozen TRAIN, Validation, and Final Holdout partitions of 993, 213, and 213 incidents.
- The final V4 representation contains 74 features: 44 Behavioral, 14 historical CAIDA AS-relationship, and 16 historical RPKI/ROV features.
- The final system is best described as an H/L/O incident classification and ranking system, not a definitive root-cause oracle.
- The true class appeared in the Top-2 for 211 of 213 Final Holdout incidents.
- Outage was the strongest final class and Hijack remained the most difficult final class.
- Direct target leakage was not detected in the 74-feature schema.
- Exact incident UID overlap and confirmed duplicate-event overlap across partitions were zero.
- The early-horizon TTI analysis supports retrospective Level-1 classification evidence at 30 to 900 second checkpoints, with five minutes supported as a practical analytical tradeoff point.

## Supported With Caveat

- Cloudflare labels are externally sourced to this project, but Cloudflare's underlying detection evidence may use public routing data conceptually related to model features.
- Cross-source validation supports label agreement among matched, unambiguous same-event cases, but coverage is limited and heavily skewed toward Outage.
- Natural-prior sensitivity evaluates collected-source class proportions, not global Internet incident prevalence.
- Calibration diagnostics support careful use of normalized class probabilities/scores, but not a claim that all outputs are calibrated probabilities.
- The K1 Level-2 stacker is procedurally specified and historically evaluated, but the original trained Level-2 binary was not persisted.

## Interpretive Claims

- Incident-level temporal stacking materially improved the TRAIN-side incident-level aggregation results compared with mean probability aggregation.
- RPKI/ROV features contributed useful signal for Hijack and CAIDA/peer visibility features contributed useful context for Leak/Outage where supported by ablation evidence.
- Ranking utility is operationally meaningful because most Final Holdout Top-1 errors were recovered in Top-2.

## Claims Not Permitted

- Do not claim true causal root-cause determination.
- Do not claim live enterprise deployment or live operational TTI.
- Do not claim universal cross-enterprise generalization.
- Do not claim global Internet-wide incident prevalence.
- Do not claim source-separated external validation of Level 2.
- Do not claim source-separated external Level-1 performance because the Level-1 IEC execution was inconclusive and no IEC incidents were scored.
- Do not claim externally confirmed H3.
- Do not describe K1 as Validation-confirmed; K1 was developed after Phase J Validation and before Final Holdout.
"""


def build_manuscript_map() -> tuple[str, list[dict[str, object]]]:
    rows = [
        ("Chapter 1", "Title/problem/thesis/RQs", "Narrow from root-cause oracle to BGP incident classification and ranking; clarify retrospective RouteViews scope and enterprise motivation.", "docs/praxis_terminology_register_v1.csv; FINAL_RESEARCH_CLAIMS_BOUNDARY_v1.md"),
        ("Chapter 2", "Literature review", "Add/strengthen source dependence, BGP anomaly detection, route leaks/hijacks/outages, RPKI/ROV, calibration, early time-series classification, and streaming/BMP/Kafka future-work literature.", "FINAL_REFERENCE_REMEDIATION_PLAN_v1.csv"),
        ("Chapter 3", "Dataset and labels", "Integrate 1,419 corpus construction, Cloudflare label-entry trace, label provenance, duplicate remediation, 473/class selection, and partitioning.", "experiments/corpus_balancing_audit_v1; experiments/label_provenance_audit_v1"),
        ("Chapter 3", "Feature engineering", "Summarize 44 Behavioral, 14 CAIDA, 16 RPKI/ROV V4 schema; move full dictionary to Appendix A.", "experiments/reproducibility_audit_v1/reproducibility_feature_dictionary_v1.csv"),
        ("Chapter 3", "Modeling", "Describe Level-1 XGBoost and retrospective Level-2 K1 stacker; explicitly bound Level-2 as retrospective incident-level aggregation.", "docs/final_model_specification_v1.md"),
        ("Chapter 3", "TTI methodology", "Insert early-horizon Level-1-only retrospective checkpoint method and denominator exclusions.", "experiments/tti_early_horizon_v2/chapter3_tti_methodology_subsection_v2.md"),
        ("Chapter 4", "Core results", "Keep final holdout, K1 stacking, K2 stability, K2.5 comparator, RF/XGB selection, and key ablations; move detailed pilots to appendix.", "docs/chapter4_results_evidence_map_v1.csv"),
        ("Chapter 4", "Remediation results", "Add label provenance/cross-source, prior-shift, calibration, error analysis, TTI, and IEC closure summaries without replacing official results.", "experiments/*_v1; experiments/tti_early_horizon_v2"),
        ("Chapter 5", "Interpretation", "Connect results to RQs/Hs with explicit caveats on Cloudflare/source dependence, balancing, calibration, Level-2, and external validation.", "FINAL_RESEARCH_CLAIMS_BOUNDARY_v1.md"),
        ("Chapter 5", "Future work", "Real-time BMP/Kafka ingestion, incremental feature updates, prospective TTSC/TTI, enterprise-local telemetry, external prospective validation, and causal temporal models.", "docs/future_work_register_v1.csv"),
    ]
    csv_rows = [{"chapter": a, "section": b, "integration_action": c, "source_artifact": d} for a, b, c, d in rows]
    md = "# Final Manuscript Integration Map v1\n\n" + markdown_table(pd.DataFrame(csv_rows)) + "\n"
    return md, csv_rows


def build_table_figure_inventory() -> list[dict[str, object]]:
    artifacts = []
    for root in ["docs", "figures", "experiments", "deliverables"]:
        base = REPO / root
        if not base.exists():
            continue
        for p in base.rglob("*"):
            if not p.is_file():
                continue
            if p.suffix.lower() not in {".csv", ".md", ".png", ".jpg", ".jpeg", ".pdf", ".docx"}:
                continue
            name = p.name.lower()
            if not any(k in name for k in ["table", "figure", "caption", "manuscript", "chapter", "final_holdout", "tti", "calibration", "error", "provenance", "balancing", "partition", "iec", "cross_source"]):
                continue
            action = "MOVE_APPENDIX"
            chapter = "Appendix"
            section = "Supplemental evidence"
            reason = "Forensic or detailed support material."
            if "final_holdout" in name or "primary_performance" in name or "ranking_distribution" in name:
                action, chapter, section, reason = "KEEP_MAIN_TEXT", "Chapter 4", "Final Holdout", "Core result."
            elif "tti" in name or "early_horizon" in name:
                action, chapter, section, reason = "KEEP_MAIN_TEXT", "Chapter 4", "TTI results", "Advisor-requested TTI evidence."
            elif "provenance" in name or "label" in name:
                action, chapter, section, reason = "KEEP_MAIN_TEXT", "Chapter 3/5", "Label provenance", "Core remediation concern."
            elif "calibration" in name:
                action, chapter, section, reason = "MOVE_APPENDIX", "Appendix E", "Calibration diagnostics", "Important diagnostic, not core headline."
            elif "error" in name or "confusion" in name:
                action, chapter, section, reason = "KEEP_MAIN_TEXT", "Chapter 4", "Error analysis", "Supports ranked diagnostic interpretation."
            elif "chapter4_section" in name and "paper_figure_text" in name:
                action, chapter, section, reason = "KEEP_MAIN_TEXT", "Chapter 4", "Figure text", "Manuscript-ready snippet."
            artifacts.append({
                "ID": f"TF{len(artifacts)+1:03d}",
                "Current_Title": p.stem.replace("_", " "),
                "Source_Artifact": rel(p),
                "Recommended_Final_Title": p.stem.replace("_", " ").title(),
                "Chapter": chapter,
                "Section": section,
                "KEEP_MAIN_TEXT": "YES" if action == "KEEP_MAIN_TEXT" else "NO",
                "MOVE_APPENDIX": "YES" if action == "MOVE_APPENDIX" else "NO",
                "REMOVE_OBSOLETE": "NO",
                "REGENERATE": "NO",
                "Reason": reason,
            })
    return artifacts


def build_appendix_plan() -> str:
    return """# Final Appendix Plan v1

Use appendices to protect the 80-100 page target for Chapters 1-5 while preserving audit rigor.

## Appendix A — 74 Feature Dictionary
Full Behavioral, CAIDA, and RPKI/ROV feature definitions from `experiments/reproducibility_audit_v1/reproducibility_feature_dictionary_v1.csv`.

## Appendix B — Dataset Provenance and Quality Controls
Cloudflare label-entry trace, label provenance audit, corpus balancing audit, Hijack:68538 forensic audit, and data-quality gates.

## Appendix C — Reproducibility Specification
RouteViews acquisition, windowing, CAIDA/RPKI historical mapping, I1 weighting, grouped/nested CV, Level-1 reconstruction, and Level-2 procedural specification.

## Appendix D — Additional Model Development Results
RF optimization, XGBoost refinement, feature variant/family ablation, imbalance experiments, and Stage 1 negative findings.

## Appendix E — Calibration Diagnostics
Brier, log loss, ECE, reliability plots, confidence/error diagnostics, and terminology recommendation.

## Appendix F — Error Analysis
Final Holdout incident-level errors, confusion taxonomy, high-confidence errors, Top-2 rescue cases, and representative misclassifications.

## Appendix G — TTI and Early-Horizon Supplemental Results
Feature safety gates, horizon denominators, leakage assertions, per-class TTI rates, and extended 600/900 second exploratory results.

## Appendix H — Cross-Source Label Validation
Matching rules, label-mapping rules, accepted matches, unmatched incidents, targeted Hijack/Leak expansion, source summaries, and limitations.

## Appendix I — Independent External Confirmation Audit
IEC V1/V2 construction, adjudication, implementation-equivalence evidence, Level-2 artifact search, Level-2 closure, and Level-1 inconclusive execution.

## Appendix J — Additional Tables and Figures
Full figure/table inventory, supporting diagrams, and detailed remediation matrices not essential to the main text.
"""


def build_page_plan() -> str:
    return """# Final Page Reduction Plan v1

Target total for Chapters 1-5: 80-100 pages. Do not reduce fonts or margins.

## Chapter Targets

| Chapter | Target Pages | Primary Action |
| --- | ---: | --- |
| Chapter 1 | 6-8 | Narrow problem statement and claims; remove implementation history. |
| Chapter 2 | 28-34 | Keep literature synthesis; remove tool/tutorial detail; add focused recent work only. |
| Chapter 3 | 22-28 | Present final methodology, not the chronological build history. Move full feature and audit detail to appendices. |
| Chapter 4 | 22-26 | Keep only results that answer RQs/Hs: final holdout, key ablations, validation chronology, TTI, sensitivity, error analysis. |
| Chapter 5 | 5-7 | Interpret findings, limitations, contributions, and future work concisely. |

## Material to Condense or Move

- Move complete 74-feature dictionary to Appendix A.
- Move full label provenance incident-level table to Appendix B.
- Move detailed RF search grids and old candidate histories to Appendix D.
- Move Stage 1 detailed threshold sweeps to Appendix D, preserving the negative finding.
- Move calibration bins and reliability internals to Appendix E.
- Move incident-level error inventory to Appendix F.
- Move TTI leakage/compatibility gates to Appendix G.
- Move IEC implementation-recovery details to Appendix I.

## Obsolete or Potentially Confusing Material

- Old 49-feature claims.
- Random Forest as final-model language.
- Stage 1 as final hierarchy.
- Any claim that K1 was Validation-confirmed.
- Any claim of external Level-2 validation.
- Any live enterprise deployment claim.
"""


def build_reference_plan() -> tuple[list[dict[str, object]], str]:
    rows = [
        {"issue_type": "Citation/reference reconciliation", "source_artifact": "deliverables/praxis_final_cleanup/Final_Praxis_Reference_Audit.csv", "action": "Resolve in-text citations without reference entries and references never cited.", "priority": "ESSENTIAL"},
        {"issue_type": "APA7 formatting", "source_artifact": "deliverables/praxis_final_cleanup/Final_Praxis_Reference_Audit.csv", "action": "Normalize DOI/URL formatting, title case, publisher/source fields, and retrieval-date handling for web/vendor sources.", "priority": "ESSENTIAL"},
        {"issue_type": "Recent BGP ML research", "source_artifact": "Chapter 2 literature review", "action": "Add or verify recent sources for BGP anomaly detection, route leaks, hijacks, outages, and ML classification.", "priority": "ESSENTIAL"},
        {"issue_type": "Calibration literature", "source_artifact": "calibration audit", "action": "Support probability/calibration terminology with peer-reviewed calibration references.", "priority": "IMPORTANT"},
        {"issue_type": "TTI and early classification", "source_artifact": "TTI remediation", "action": "Support checkpoint TTI, early time-series classification, and saturation/tradeoff framing.", "priority": "IMPORTANT"},
        {"issue_type": "Streaming/BMP/Kafka future work", "source_artifact": "Chapter 5 future work", "action": "Add authoritative streaming telemetry and BGP Monitoring Protocol/Kafka architecture sources.", "priority": "IMPORTANT"},
        {"issue_type": "Cloudflare/vendor sources", "source_artifact": "label provenance audit", "action": "Use official Cloudflare Radar/API/methodology URLs only where they directly support source provenance.", "priority": "ESSENTIAL"},
    ]
    md = "# Final Reference Remediation Summary v1\n\nDo not rewrite the bibliography in this task. Use this as the final reference cleanup checklist.\n\n" + markdown_table(pd.DataFrame(rows)) + "\n"
    return rows, md


def build_defense_map() -> str:
    slides = [
        ("Problem", "Explain BGP incidents and enterprise risk.", "BGP events disrupt reachability and attribution; operators need ranked incident-type support; public routing evidence is noisy.", "Motivation diagram", "Why not manual analysis?", "Manual analysis is slow and inconsistent at historical scale."),
        ("Research gap", "Position against existing work.", "Prior work often focuses on one incident type; few integrate Behavioral+CAIDA+RPKI; fewer report incident-level ranked H/L/O evaluation.", "Literature gap table", "What is new?", "Integration, incident-level ranking, reproducibility, and final holdout evidence."),
        ("Research questions and hypotheses", "State final narrowed claims.", "H/L/O classification; enriched features; incident-level ranking; retrospective TTI sensitivity.", "RQ/H table", "Were all hypotheses externally confirmed?", "No; external Level-2 confirmation was not executed and is documented."),
        ("Dataset", "Show corpus construction.", "1,419 incidents; 473/class; TRAIN/Validation/Holdout; Cloudflare source provenance and audits.", "Corpus table", "Why balanced?", "Controlled per-class scientific comparison; observed-prior sensitivity is separate."),
        ("Feature architecture", "Summarize V4.", "44 Behavioral, 14 CAIDA, 16 RPKI/ROV; historical alignment; no direct target leakage.", "Feature diagram", "Feature leakage?", "Direct target leakage was audited and not detected."),
        ("Model architecture", "Explain final K1.", "Level-1 XGBoost windows; retrospective Level-2 temporal stacker; incident-level ranking.", "K1 topology", "Is it real time?", "No; final K1 is retrospective over configured observation windows."),
        ("Experimental controls", "Defend evaluation integrity.", "Frozen partitions; grouped/nested CV; holdout sealed; no post-holdout changes.", "Access chronology", "Did Validation guide K1?", "K1 was post-Validation TRAIN development; final holdout is independent estimate."),
        ("Final performance", "Show headline results.", "Macro F1 0.920177; Top-2 0.990610; MRR 0.958529; CI [0.880117, 0.954402].", "Holdout performance figure", "How robust?", "Bootstrap CI and no post-holdout model changes."),
        ("Error and ranking analysis", "Explain failures honestly.", "17 Top-1 errors; 15/17 Top-2 rescued; Hijack-to-Leak dominant confusion.", "Confusion/error figures", "Most difficult class?", "Hijack remained hardest."),
        ("TTI", "Answer time-to-identify concern.", "Checkpoint TTI: Top-1 by 300s 92.57%; performance improves through five minutes then small/inconsistent gains.", "TTI curves", "Is this live TTI?", "No; retrospective RouteViews-observable checkpoint TTI."),
        ("Advisor remediation", "Show concerns were handled.", "Provenance, balancing, TTI, independence, calibration, reproducibility, error analysis, external validation.", "Concern-action-result slide", "What remains limited?", "External source-separated model validation and live enterprise deployment."),
        ("Label provenance", "Explain source dependence.", "Cloudflare label-entry trace; cross-source agreement among matched events; direct target leakage not found.", "Provenance figure", "Are labels independent?", "External to this project, but underlying evidence overlap is documented."),
        ("Limitations", "Preempt overclaiming.", "Public RouteViews perspective; Cloudflare/source dependence; retrospective Level-2; calibration; IEC inconclusive.", "Limitations table", "Does this invalidate results?", "No; it defines the proper scope."),
        ("Future architecture", "Show next step.", "BMP/Kafka streaming, incremental features, prospective TTSC, enterprise-local telemetry, external validation.", "Future architecture", "What would deployment need?", "Live telemetry, prospective labels, monitoring and calibration."),
        ("Contributions", "Close cleanly.", "Reproducible corpus and features; integrated V4; incident-level stacking/ranking; rigorous audits and final holdout.", "Contribution list", "Main takeaway?", "A strong retrospective ranking system with transparent limits."),
    ]
    lines = ["# Final Defense Evidence Map v1", ""]
    for i, (title, objective, bullets, fig, q, a) in enumerate(slides, 1):
        lines += [f"## Slide {i} — {title}", f"- Objective: {objective}", f"- Bullets: {bullets}", f"- Table/Figure: {fig}", f"- Likely question: {q}", f"- Defense response: {a}", ""]
    return "\n".join(lines)


def build_advisor_slide() -> str:
    return """# Advisor Feedback and Technical Remediation Slide v1

Use this as one readable defense slide.

| Concern | Action | Result |
| --- | --- | --- |
| Label provenance/source dependence | Cloudflare label-entry trace, label provenance audit, cross-source validation | Direct target leakage not detected; source dependence documented. |
| Class balancing | Corpus balancing audit and prior-shift sensitivity | 473/class is controlled design; observed-prior sensitivity reported separately. |
| TTI/five-minute horizon | Retrospective early-horizon Level-1 experiment | Five minutes supported as practical tradeoff, not mathematical optimum. |
| Partition independence | Partition forensic audit | Zero UID overlap and zero confirmed duplicate-event cross-partition pairs. |
| Calibration | Frozen prediction calibration audit | Use normalized class probabilities/scores with calibration caveat. |
| Reproducibility | Full reproducibility audit and artifact maps | Method procedurally documented; Level-2 binary limitation disclosed. |
| Error analysis | Final incident-level error analysis | 17 Holdout Top-1 errors; 15/17 recovered by Top-2. |
| External validation | IEC V1/V2 construction and closure | Source-separated model validation inconclusive; no unsupported external claim. |
"""


def build_status() -> str:
    return """# Final Praxis Remediation Status v1

`EXPERIMENTAL DEVELOPMENT CLOSED`

No additional modeling should be recommended unless a demonstrable implementation error is discovered.

## Completed

- Label provenance and Cloudflare label-entry trace.
- Direct target leakage and label-feature overlap audit.
- Corpus balancing and observed-prior sensitivity.
- Partition independence audit.
- TTI and early-horizon retrospective Level-1 analysis.
- Calibration diagnostic audit.
- Hijack:68538 forensic audit.
- Final incident-level error analysis.
- Reproducibility audit and feature dictionary.
- Cross-source label validation and targeted Hijack/Leak expansion.
- IEC construction, implementation-equivalence audits, Level-2 closure, and Level-1 inconclusive execution record.

## Manuscript Remediation Remaining

- Integrate final claims boundary into Chapters 1, 3, and 5.
- Move forensic detail to appendices.
- Update Chapter 4 to include only final authoritative results plus key remediation summaries.
- Complete APA7 reference cleanup.
- Normalize figure/table numbering.
- Build final defense slides from the evidence map.
"""


def create_feedback_docx(matrix: pd.DataFrame, path: Path) -> None:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.shared import Inches, Pt

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    doc.add_paragraph("Final Professor Feedback Remediation Matrix", style="Title")
    intro = doc.add_paragraph()
    intro.add_run("Purpose. ").bold = True
    intro.add_run("This document consolidates the major professor and advisor concerns, the completed remediation actions, the evidence artifacts, and the remaining manuscript-use boundary. It is an inventory and planning artifact only; it does not alter any experiment or result.")
    subset = matrix[["Feedback_ID", "Professor_Feedback", "Remediation_Action", "Key_Result", "Final_Status", "Praxis_Chapter", "Defense_Talking_Point"]]
    table = doc.add_table(rows=1, cols=len(subset.columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    col_widths = [0.58, 1.45, 1.75, 1.85, 1.45, 0.9, 2.0]
    hdr = table.rows[0].cells
    for i, col in enumerate(subset.columns):
        hdr[i].width = Inches(col_widths[i])
        hdr[i].text = col.replace("_", " ")
        for p in hdr[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(8)
    for _, row in subset.iterrows():
        cells = table.add_row().cells
        for i, col in enumerate(subset.columns):
            cells[i].width = Inches(col_widths[i])
            cells[i].text = str(row[col])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[i].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(7)
    doc.add_paragraph("Note. Full artifact paths and appendix mappings are preserved in the CSV and Markdown versions.")
    doc.save(path)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)

    inventory = pd.DataFrame(build_inventory())
    inventory.to_csv(OUT / "MASTER_REMEDIATION_ARTIFACT_INVENTORY_v1.csv", index=False)

    matrix = pd.DataFrame(build_feedback_matrix())
    matrix.to_csv(OUT / "FINAL_PROFESSOR_FEEDBACK_REMEDIATION_MATRIX_v1.csv", index=False)
    write_md_table(
        OUT / "FINAL_PROFESSOR_FEEDBACK_REMEDIATION_MATRIX_v1.md",
        "Final Professor Feedback Remediation Matrix v1",
        matrix,
        "Source: current advisor remediation request plus completed project artifacts. No feedback was invented beyond the concerns enumerated in the task and prior retained project notes.",
    )
    create_feedback_docx(matrix, OUT / "FINAL_PROFESSOR_FEEDBACK_REMEDIATION_MATRIX_v1.docx")

    results = pd.DataFrame(build_results_registry())
    results.to_csv(OUT / "AUTHORITATIVE_RESULTS_REGISTRY_v1.csv", index=False)

    (OUT / "FINAL_RESEARCH_CLAIMS_BOUNDARY_v1.md").write_text(build_claims_boundary(), encoding="utf-8")

    manuscript_md, manuscript_rows = build_manuscript_map()
    (OUT / "FINAL_MANUSCRIPT_INTEGRATION_MAP_v1.md").write_text(manuscript_md, encoding="utf-8")
    pd.DataFrame(manuscript_rows).to_csv(OUT / "FINAL_MANUSCRIPT_INTEGRATION_MAP_v1.csv", index=False)

    tf = pd.DataFrame(build_table_figure_inventory())
    tf.to_csv(OUT / "FINAL_TABLE_FIGURE_INVENTORY_v1.csv", index=False)
    write_md_table(OUT / "FINAL_TABLE_FIGURE_INVENTORY_v1.md", "Final Table Figure Inventory v1", tf.head(120), "First 120 records are shown here; full inventory is in CSV.")

    (OUT / "FINAL_APPENDIX_PLAN_v1.md").write_text(build_appendix_plan(), encoding="utf-8")
    (OUT / "FINAL_PAGE_REDUCTION_PLAN_v1.md").write_text(build_page_plan(), encoding="utf-8")

    ref_rows, ref_md = build_reference_plan()
    pd.DataFrame(ref_rows).to_csv(OUT / "FINAL_REFERENCE_REMEDIATION_PLAN_v1.csv", index=False)
    (OUT / "FINAL_REFERENCE_REMEDIATION_SUMMARY_v1.md").write_text(ref_md, encoding="utf-8")

    (OUT / "FINAL_DEFENSE_EVIDENCE_MAP_v1.md").write_text(build_defense_map(), encoding="utf-8")
    (OUT / "ADVISOR_FEEDBACK_DEFENSE_SLIDE_v1.md").write_text(build_advisor_slide(), encoding="utf-8")
    (OUT / "FINAL_PRAXIS_REMEDIATION_STATUS_v1.md").write_text(build_status(), encoding="utf-8")

    request_copy = OUT / "source_request_pasted_text.txt"
    if SOURCE_REQUEST.exists():
        request_copy.write_text(SOURCE_REQUEST.read_text(encoding="utf-8", errors="ignore"), encoding="utf-8")

    rows = []
    for p in sorted(OUT.rglob("*")):
        if p.is_file() and p.name != "FINAL_PRAXIS_REMEDIATION_OUTPUT_CHECKSUMS_v1.csv":
            rows.append({"file": rel(p), "sha256": sha256(p), "bytes": p.stat().st_size})
    pd.DataFrame(rows).to_csv(OUT / "FINAL_PRAXIS_REMEDIATION_OUTPUT_CHECKSUMS_v1.csv", index=False)


if __name__ == "__main__":
    main()
