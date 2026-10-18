from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


OUT_DIR = Path("/Users/biruk/Documents/New project/deliverables/advising_reports")
OUT_DOCX = OUT_DIR / "Advising_Report_2026-09-06.docx"


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def add_paragraph(doc, text="", style=None, bold=False, align=None):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)
    if align is not None:
        p.alignment = align
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    return table


def configure_styles(doc):
    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(11)
    for name, size in [("Title", 16), ("Heading 1", 13), ("Heading 2", 12)]:
        styles[name].font.name = "Times New Roman"
        styles[name].font.size = Pt(size)
        styles[name].font.bold = True


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    doc = Document()
    configure_styles(doc)

    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Doctoral Praxis Progress Report")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(16)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Advising Session: September 6, 2026")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)

    meta_rows = [
        ("Student", "Biruk D. Mekonnen"),
        ("Program", "Doctor of Engineering in Artificial Intelligence"),
        ("Praxis Title", "An Artificial Intelligence-Based Classification and Ranking System for Border Gateway Protocol Incidents in Enterprise Networks"),
        ("Report Purpose", "Summarize completed methodological, experimental, and writing progress since the August 23 advising report and identify the remaining items before final submission/defense preparation."),
    ]
    add_table(doc, ["Item", "Current Value"], meta_rows)

    doc.add_heading("1. Executive Summary", level=1)
    add_paragraph(
        doc,
        "Since the August 23 advising report, the project moved from a corrected experimental design to a completed, leakage-controlled Praxis pipeline with final independent holdout results. The incident roster was rebuilt and frozen, the RouteViews feature pipeline was completed, Behavioral, CAIDA, and RPKI/ROV features were finalized, and the final K1 incident-level stacking model was evaluated once on the sealed Final Holdout.",
    )
    add_paragraph(
        doc,
        "The current work is no longer primarily an experimentation problem. The remaining work is dissertation construction: ensuring that Chapters 3 and 4 describe the final methods/results cleanly, that Chapter 5 interprets the findings with appropriate limitations, and that all tables, figures, citations, and front-matter elements are consistent with the final evidence.",
    )

    doc.add_heading("2. Progress Since the August 23 Advising Report", level=1)
    add_bullets(
        doc,
        [
            "Rebuilt and froze the incident-level corpus as Roster v2 with 1,419 independent incidents: 473 Hijack, 473 Leak, and 473 Outage.",
            "Created leakage-safe incident-level partitions: TRAIN = 993 incidents, Validation = 213 incidents, and Final Holdout = 213 incidents.",
            "Completed RouteViews-based feature construction for the TRAIN corpus and preserved the incident as the independent experimental unit.",
            "Finalized the V4 feature representation with 74 features: 44 Behavioral, 14 CAIDA AS-relationship features, and 16 historical RPKI/ROV features.",
            "Evaluated baseline Random Forest, tuned Random Forest, XGBoost, refined XGBoost, imbalance strategies, feature variants, and feature-family ablations using TRAIN-only grouped cross-validation.",
            "Completed the Phase J one-time Validation confirmation for the refined XGBoost candidate without using Validation for further tuning.",
            "Developed K1 as a post-Validation TRAIN-only incident-level stacking architecture and completed leakage/stability checks before Final Holdout access.",
            "Completed the one-time Final Holdout evaluation after the deployable K1 model was frozen.",
            "Generated Chapter 4-ready figures and tables for imbalance, Stage 1, Random Forest, XGBoost, Validation, K1 stacking, K2 stability, Final Holdout, error analysis, and latency.",
            "Produced a condensed final Praxis draft with page-count compliance across the overall document and most chapters.",
        ],
    )

    doc.add_heading("3. Current Authoritative Experimental State", level=1)
    add_table(
        doc,
        ["Component", "Authoritative Status"],
        [
            ("Incident corpus", "Frozen Roster v2: 1,419 independent incidents"),
            ("Class balance", "473 Hijack / 473 Leak / 473 Outage"),
            ("Experimental split", "TRAIN 993 / Validation 213 / Final Holdout 213"),
            ("Feature schema", "V4: 74 total features"),
            ("Feature families", "44 Behavioral + 14 CAIDA + 16 RPKI/ROV"),
            ("TRAIN Behavioral rows", "68,031 five-minute incident windows"),
            ("Final model", "K1 XGB-to-XGB incident-level temporal stacking architecture"),
            ("Stage 1 status", "Experimental/auxiliary; not used as the final deployed classifier"),
            ("Evaluation integrity", "Validation opened once in Phase J; Final Holdout opened once in Phase L; no post-Holdout model changes"),
        ],
    )

    doc.add_heading("4. Final Holdout Results", level=1)
    add_table(
        doc,
        ["Metric", "Final Holdout Value"],
        [
            ("Final Holdout incidents", "213 total; 71 Hijack / 71 Leak / 71 Outage"),
            ("Final Holdout window rows", "14,906 five-minute rows"),
            ("Incident Macro F1", "0.920177"),
            ("95% CI for Macro F1", "[0.880117, 0.954402]"),
            ("Top-1 Accuracy", "0.920188"),
            ("Top-2 Accuracy", "0.990610"),
            ("Mean Reciprocal Rank", "0.958529"),
            ("Hijack F1", "0.877698"),
            ("Leak F1", "0.904110"),
            ("Outage F1", "0.978723"),
            ("Ranking distribution", "Rank 1 = 196; Rank 2 = 15; Rank 3 = 2"),
            ("Top-2 count", "211 of 213 incidents"),
        ],
    )
    add_paragraph(
        doc,
        "The final results support the main Praxis claim that the completed system can classify and rank BGP incident categories at the incident level using RouteViews-derived evidence. Outage is the strongest final class, while Hijack remains the most difficult final class. These claims are stated as model-performance findings, not as proof of real-time production deployment.",
    )

    doc.add_heading("5. Issues Addressed From Prior Advisor Feedback", level=1)
    add_table(
        doc,
        ["Prior Feedback / Risk", "Current Resolution"],
        [
            ("Incident-label leakage and row-level leakage", "Partitions, grouped CV, and final evaluation use incident identity as the independent unit."),
            ("Duplicated or correlated incident records", "Roster v2 duplicate remediation completed before freezing the final corpus."),
            ("Unclear incident windows", "Timestamp semantics were reviewed; Hijack/Leak active-window handling was separated from silent detector gaps; outage intervals use documented start/end times."),
            ("Need for final independent counts by class and partition", "Counts are frozen and documented: 473 per class, 993 TRAIN, 213 Validation, 213 Final Holdout."),
            ("Need for macro-F1, per-class F1, ranking metrics, and confidence intervals", "All are now available for Final Holdout and mapped to source artifacts."),
            ("Need for ablations and baselines", "Feature-variant, feature-family, RF, XGBoost, and stacking comparisons have been documented."),
            ("Risk of overstating RouteViews as enterprise-local telemetry", "The final writing frames RouteViews as public interdomain routing evidence and lists observability limitations."),
        ],
    )

    doc.add_heading("6. Dissertation Writing and Formatting Status", level=1)
    add_table(
        doc,
        ["Area", "Status"],
        [
            ("Overall draft", "Condensed final Praxis draft generated; current rendered length is approximately 93 pages."),
            ("Chapter 1", "Within preferred range after cleanup."),
            ("Chapter 2", "Still comparatively short relative to the 30-35 page best-practice guidance; may require advisor decision on whether depth is sufficient."),
            ("Chapter 3", "Updated to final methodology, including V4 features, partitions, leakage controls, and K1 architecture."),
            ("Chapter 4", "Updated with final results tables/figures and experimental progression."),
            ("Chapter 5", "Prepared around findings, limitations, future work, and operational interpretation."),
            ("Front matter", "Contains placeholders where official committee/front-matter wording or personal sections still need final user-provided text."),
            ("References", "APA 7 formatting has been improved; final source-by-source verification should still be completed before submission."),
        ],
    )

    doc.add_heading("7. Items Recommended for Advisor Review on September 6", level=1)
    add_bullets(
        doc,
        [
            "Confirm that the final title should use “classification and ranking” rather than “root-cause ranking” as the primary framing.",
            "Confirm whether Chapter 2 depth is acceptable or whether additional peer-reviewed literature synthesis should be added.",
            "Confirm that Stage 1 should remain explicitly auxiliary/unresolved rather than part of the final deployed architecture.",
            "Confirm that K1 should be described as post-window incident classification/ranking, with real-time incremental inference left as future work.",
            "Confirm whether detailed RouteViews/MRT engineering pilots should remain in appendices rather than the main Chapter 4 narrative.",
            "Confirm final committee names, institutional certification language, dedication, acknowledgements, and any required front-matter wording.",
        ],
    )

    doc.add_page_break()
    doc.add_heading("8. Remaining Work Before Final Submission", level=1)
    add_table(
        doc,
        ["Priority", "Task", "Reason"],
        [
            ("High", "Finalize official front matter and committee information", "Required for submission-ready Praxis formatting."),
            ("High", "Complete a final APA 7 reference audit", "Prevents citation/reference mismatch and style errors."),
            ("High", "Review Chapter 2 length and literature coverage with advisor", "Only major chapter-length item currently marked for review."),
            ("High", "Lock figure/table numbering after final document structure is fixed", "Avoids broken cross-references during final edits."),
            ("Medium", "Prepare defense slide deck from the final evidence map", "Supports committee presentation and oral defense readiness."),
            ("Medium", "Finalize Chapter 5 future-work language on live/streaming deployment", "Keeps performance claims accurate and defensible."),
            ("Medium", "Decide whether optional SHAP/explainability figures are needed", "Useful, but should not delay the final submission if not required."),
        ],
    )

    doc.add_heading("9. Proposed Advisor Discussion Questions", level=1)
    add_bullets(
        doc,
        [
            "Does the current final result provide sufficient evidence for the Praxis contribution?",
            "Should Chapter 2 be expanded, or is the current literature review acceptable for the final submission timeline?",
            "Is the distinction between Validation, post-Validation K1 TRAIN-only development, and Final Holdout evaluation clear enough?",
            "Should the final document emphasize Top-2 ranking performance as the most operationally useful outcome?",
            "Which remaining writing items should be completed before circulating the next full draft?",
        ],
    )

    doc.add_heading("10. Attachments / Supporting Artifacts Available", level=1)
    add_bullets(
        doc,
        [
            "Final_Praxis_Submission.docx and Final_Praxis_Submission.pdf",
            "Master Praxis experimental record and authoritative-number register",
            "Chapter 3 methodology pipeline figure options",
            "Chapter 4 results tables and high-resolution figures",
            "Final Holdout metrics, ranking distribution, confusion/error analysis, and latency documentation",
            "Evaluation-integrity narrative documenting TRAIN, Validation, and Final Holdout access chronology",
        ],
    )

    add_paragraph(doc, "")
    add_paragraph(
        doc,
        "Prepared for the September 6, 2026 advising discussion.",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    main()
