from __future__ import annotations

import csv
import re
import shutil
import subprocess
import tempfile
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path("/Users/biruk/Documents/New project")
TEMPLATE = Path("/Users/biruk/Desktop/Advising session/08_09_2026/Biruk Mekonnen.docx")
OUT_DIR = ROOT / "deliverables" / "praxis"
DOCX_OUT = OUT_DIR / "Biruk_Mekonnen_Final_Praxis_APA7_v1.docx"
PDF_OUT = OUT_DIR / "Biruk_Mekonnen_Final_Praxis_APA7_v1.pdf"
SOFFICE = Path(
    "/Users/biruk/.cache/codex-runtimes/codex-primary-runtime/dependencies/"
    "native/libreoffice-headless/libreoffice/LibreOfficeDev.app/Contents/MacOS/soffice"
)


TITLE = (
    "An Artificial Intelligence-Based Classification and Ranking System for "
    "Border Gateway Protocol Incidents in Enterprise Networks"
)
AUTHOR = "Biruk D. Mekonnen"
SUBMISSION_DATE = "August 15, 2026"


def clear_document(doc: Document) -> None:
    body = doc._body._element
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder or "Update field in Word if needed"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(text)
    run._r.append(fld_end)


def add_page_break(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def p(doc: Document, text: str, style: str = "Normal", bold: bool = False, italic: bool = False):
    paragraph = doc.add_paragraph(style=style)
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    return paragraph


def h(doc: Document, text: str, level: int = 1):
    return doc.add_heading(text, level=level)


def add_caption(doc: Document, text: str):
    try:
        return p(doc, text, style="Caption")
    except Exception:
        return p(doc, text, italic=True)


def add_table(doc: Document, title: str, headers: list[str], rows: list[list[str]]) -> None:
    add_caption(doc, title)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for idx, header in enumerate(headers):
        hdr[idx].text = header
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = str(value)
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)
    doc.add_paragraph()


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def current_references() -> list[str]:
    doc = Document(TEMPLATE)
    refs: list[str] = []
    start = None
    for idx, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip().lower() == "references":
            start = idx
    if start is None:
        return refs
    for paragraph in doc.paragraphs[start + 1 :]:
        text = " ".join(paragraph.text.split())
        if text:
            refs.append(text)
    return refs


def source_chapter2_paragraphs() -> list[tuple[str, str]]:
    doc = Document(TEMPLATE)
    start = None
    end = None
    for idx, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if text == "Chapter 2—Literature Review":
            start = idx
        if start is not None and text.startswith("Chapter 3"):
            end = idx
            break
    if start is None or end is None:
        return []
    out: list[tuple[str, str]] = []
    for paragraph in doc.paragraphs[start + 1 : end]:
        text = " ".join(paragraph.text.split()).strip()
        if not text:
            continue
        text = text.replace("root-cause ranking system", "classification and ranking system")
        text = text.replace("root-cause analysis", "incident-cause analysis")
        text = text.replace("root-cause diagnosis", "incident-cause diagnosis")
        text = text.replace("root-cause", "incident-cause")
        text = text.replace("root cause", "incident cause")
        if "final methodology of this praxis emphasizes a hierarchical two-stage Random Forest model" in text:
            text = (
                "The final methodology of this Praxis emphasizes a leakage-safe "
                "incident-level XGBoost stacking architecture, while Random Forest is "
                "retained as an important benchmark. The literature review therefore "
                "evaluates model families according to their relevance to structured, "
                "protocol-derived BGP features, incident-level grouping, ranking output, "
                "and operational interpretability."
            )
        if text.startswith("Random Forest is particularly attractive for this praxis"):
            text = (
                "Tree-based ensemble models are attractive for this Praxis because the "
                "input is incident-level tabular data composed of behavioral, "
                "AS-relationship, and RPKI/ROV features. Random Forest provides a strong "
                "baseline, while XGBoost provides a boosted-tree alternative that proved "
                "stronger in the final controlled comparison."
            )
        if "final model used a two-stage Random Forest" in text:
            text = text.replace(
                "The final model used a two-stage Random Forest design.",
                "The final model used an incident-level XGBoost stacking design.",
            )
        out.append((paragraph.style.name, text))
    return out


def add_front_matter(doc: Document) -> None:
    for _ in range(2):
        doc.add_paragraph()
    title_para = p(doc, TITLE)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    by_para = p(doc, f"by {AUTHOR}")
    by_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    for line in [
        "B.S. in Electrical Engineering, July 2007, Mekelle University",
        "M.S. in Network Engineering, May 2023, Southern Methodist University",
        "",
        "A Praxis submitted to",
        "The Faculty of",
        "The School of Engineering and Applied Science",
        "of The George Washington University",
        "in partial fulfillment of the requirements",
        "for the degree of Doctor of Engineering",
        "",
        SUBMISSION_DATE,
        "",
        "Praxis directed by",
        "Timothy Blackburn, PhD",
        "Professorial Lecturer of Engineering Management and Systems Engineering",
        "Amir Etemadi, PhD",
        "Assistant Professor of Engineering and Applied Science",
    ]:
        para = p(doc, line)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_break(doc)
    p(
        doc,
        (
            "The School of Engineering and Applied Science of The George Washington "
            "University certifies that Biruk D. Mekonnen has submitted this Praxis "
            "for review in partial fulfillment of the requirements for the Doctor "
            "of Engineering degree. Final committee signature details should be "
            "completed using the official approved form."
        ),
    )
    doc.add_paragraph()
    p(doc, TITLE)
    p(doc, AUTHOR)
    doc.add_paragraph()
    p(doc, "Praxis Research Committee:")
    p(doc, "Dr. Amir Etemadi, Assistant Professor of Engineering and Applied Science, Praxis Director")
    p(doc, "Dr. Timothy Blackburn, Professorial Lecturer of Engineering Management and Systems Engineering, Praxis Co-director")
    p(doc, "[Third committee member name and title to be confirmed], Committee Member")
    add_page_break(doc)
    p(doc, f"Copyright {date.today().year} by Biruk Mekonnen")
    p(doc, "All rights reserved")
    add_page_break(doc)
    h(doc, "Dedication")
    p(doc, "[Dedication text to be finalized by the author.]")
    add_page_break(doc)
    h(doc, "Acknowledgements")
    p(doc, "[Acknowledgements text to be finalized by the author.]")
    add_page_break(doc)
    h(doc, "Abstract of Praxis")
    p(doc, TITLE)
    abstract = [
        (
            "Border Gateway Protocol (BGP) incidents can disrupt reachability, traffic "
            "engineering, and operational confidence across interdomain networks. "
            "Prefix hijacks, route leaks, and outages may create overlapping symptoms "
            "in public routing telemetry, making it difficult for operators to quickly "
            "prioritize the most likely incident cause. This Praxis develops and "
            "evaluates an artificial intelligence-based system that classifies and "
            "ranks three BGP incident categories using historical RouteViews UPDATE "
            "data, incident-level labels, and routing-security context."
        ),
        (
            "The final pipeline uses a frozen corpus of 1,419 independent incidents, "
            "balanced across 473 hijack, 473 leak, and 473 outage incidents. All splits "
            "were performed at the incident level: 993 incidents for TRAIN, 213 for "
            "Validation, and 213 for the untouched Final Holdout. Five-minute windows "
            "were represented with 74 frozen V4 features: 44 behavioral BGP features, "
            "14 CAIDA AS-relationship features, and 16 RPKI/ROV features."
        ),
        (
            "Model development began with baseline tree-based classifiers and progressed "
            "through imbalance treatment, feature-variant ablation, feature-family "
            "ablation, Random Forest tuning, XGBoost refinement, and a final incident-level "
            "stacking architecture. The terminal model, A0_XGB_TO_XGB K1, combines "
            "window-level XGBoost probability evidence with incident-level temporal "
            "summary features and a second-level XGBoost classifier."
        ),
        (
            "On the untouched Final Holdout of 213 incidents, the final system achieved "
            "incident-level Macro F1 of 0.920177 with a 95% confidence interval of "
            "[0.880117, 0.954402]. Top-1 accuracy was 0.920188, Top-2 accuracy was "
            "0.990610, and mean reciprocal rank was 0.958529. The true class was ranked "
            "in the Top-2 for 211 of 213 incidents. Per-class F1 scores were 0.877698 "
            "for Hijack, 0.904110 for Leak, and 0.978723 for Outage."
        ),
        (
            "The study demonstrates that combining behavioral BGP features with CAIDA "
            "relationship context and RPKI/ROV origin-validation signals can support "
            "accurate incident-level classification and ranked incident-cause output. "
            "The evidence also shows important limitations: public RouteViews collectors "
            "do not fully represent enterprise-local routing state, incident labels "
            "depend on external sources, and the Normal-vs-Incident Stage 1 formulation "
            "remained experimental rather than part of the final architecture."
        ),
    ]
    for paragraph in abstract:
        p(doc, paragraph)
    add_page_break(doc)
    h(doc, "Table of Contents")
    for entry in [
        ("Abstract of Praxis", "7"),
        ("List of Figures", "10"),
        ("List of Tables", "11"),
        ("List of Symbols", "12"),
        ("List of Acronyms", "13"),
        ("Chapter 1-Introduction", "14"),
        ("Chapter 2-Literature Review", "17"),
        ("Chapter 3-Methodology", "48"),
        ("Chapter 4-Results", "54"),
        ("Chapter 5-Conclusions", "59"),
        ("References", "62"),
    ]:
        p(doc, f"{entry[0]}    {entry[1]}")
    add_page_break(doc)
    h(doc, "List of Figures")
    for entry in [
        "Figure 3.1 Final Research Workflow",
        "Figure 3.2 Final K1 Incident-Level Stacking Architecture",
    ]:
        p(doc, entry)
    add_page_break(doc)
    h(doc, "List of Tables")
    for entry in [
        "Table 1.1 Research Questions and Evidence",
        "Table 3.1 Incident Partition Summary",
        "Table 3.2 Frozen Feature Families",
        "Table 3.3 Feature Variant Definitions",
        "Table 3.4 Modeling Phases",
        "Table 3.5 Evaluation-Integrity Controls",
        "Table 4.1 Model Performance Progression",
        "Table 4.2 Refined RF and Refined XGBoost TRAIN Grouped-CV Comparison",
        "Table 4.3 Phase J Validation Metrics",
        "Table 4.4 K1 Stability and Final Pre-Holdout Reference",
        "Table 4.5 Final Holdout Performance",
        "Table 4.6 Final Holdout Per-Class F1",
        "Table 4.7 Final Holdout Ranking Distribution",
        "Table 5.1 Research Question Findings",
    ]:
        p(doc, entry)
    add_page_break(doc)
    h(doc, "List of Symbols")
    symbols = [
        ("AS", "Autonomous System"),
        ("AS_PATH", "Ordered AS sequence in a BGP route announcement"),
        ("BGP", "Border Gateway Protocol"),
        ("F1", "Harmonic mean of precision and recall"),
        ("MRR", "Mean reciprocal rank"),
        ("NLRI", "Network Layer Reachability Information"),
        ("ROA", "Route Origin Authorization"),
        ("ROV", "Route Origin Validation"),
        ("RPKI", "Resource Public Key Infrastructure"),
    ]
    add_table(doc, "Table S.1 Key Symbols and Abbreviations", ["Symbol", "Meaning"], symbols)
    add_page_break(doc)
    h(doc, "List of Acronyms")
    for acronym, meaning in symbols:
        p(doc, f"{acronym}    {meaning}")
    add_page_break(doc)


def add_chapter1(doc: Document) -> None:
    h(doc, "Chapter 1-Introduction")
    h(doc, "Background and Research Motivation", 2)
    for paragraph in [
        (
            "The Internet depends on BGP to exchange reachability information among "
            "autonomous systems. Because BGP was designed around distributed trust and "
            "local routing policy, abnormal routing events can propagate quickly and "
            "may be difficult to interpret from control-plane evidence alone. Three "
            "incident categories are central to this Praxis: prefix hijacks, route "
            "leaks, and outages. Each can affect reachability or traffic paths, yet "
            "the observable symptoms in BGP UPDATE data may overlap."
        ),
        (
            "Enterprise network teams and network operations centers typically rely on "
            "monitoring dashboards, route collectors, alerts, and engineer judgment to "
            "interpret routing anomalies. These tools may show that routing behavior "
            "changed, but they often provide limited support for ranking the most likely "
            "incident type from routing evidence. This creates a gap between anomaly "
            "detection and operator-facing incident interpretation."
        ),
        (
            "This Praxis addresses that gap by treating incident classification and "
            "ranking as an empirical machine-learning problem. The study does not claim "
            "to infer private enterprise-local root causes from public telemetry. Instead, "
            "it asks whether a rigorously labeled, leakage-controlled corpus of public "
            "BGP observations can support reliable classification and ranked incident-cause "
            "output for hijack, leak, and outage incidents."
        ),
    ]:
        p(doc, paragraph)
    h(doc, "Problem Statement", 2)
    p(
        doc,
        (
            "BGP incidents can be operationally disruptive, but their symptoms are often "
            "ambiguous in public routing telemetry. Existing approaches frequently focus "
            "on detecting anomalies or on one specialized incident type. Operators need "
            "a reproducible method that combines BGP behavioral evidence, AS-relationship "
            "context, and route-origin validation signals to classify and rank likely "
            "incident categories without leaking future evaluation information."
        ),
    )
    h(doc, "Thesis Statement", 2)
    p(
        doc,
        (
            "An incident-level machine-learning system that combines BGP behavioral "
            "features, CAIDA AS-relationship context, and RPKI/ROV route-origin validation "
            "signals can accurately classify and rank hijack, leak, and outage incidents "
            "when trained and evaluated under leakage-safe incident-level partitions."
        ),
    )
    h(doc, "Research Objectives", 2)
    objectives = [
        "Construct a balanced, quality-controlled, incident-level BGP corpus for hijack, leak, and outage events.",
        "Extract reproducible five-minute behavioral features from historical RouteViews UPDATE data.",
        "Enrich the behavioral representation with CAIDA AS-relationship and historical RPKI/ROV features.",
        "Evaluate model families and feature variants using incident-grouped validation to prevent leakage across windows from the same incident.",
        "Develop and freeze a final incident-level classification and ranking model before a one-time Final Holdout evaluation.",
        "Quantify classification, ranking, uncertainty, limitations, and future operational-readiness requirements.",
    ]
    for item in objectives:
        p(doc, item, style="List Bullet")
    h(doc, "Research Questions and Hypotheses", 2)
    rq_rows = [
        ["RQ1", "Can historical RouteViews data support a reproducible BGP incident feature corpus?", "Supported by a frozen TRAIN behavioral corpus of 68,031 five-minute rows and a V4 feature schema."],
        ["RQ2", "Do CAIDA and RPKI enrichments improve incident classification?", "Supported with caveat by V1-V4 feature variants and feature-family ablation."],
        ["RQ3", "Which architecture best classifies and ranks incident causes?", "Supported by TRAIN grouped-CV comparisons, K1 stacking, K2 stability checks, and K2.5 architecture comparison."],
        ["RQ4", "Does the locked final system generalize to unseen incidents?", "Supported by one-time Final Holdout evaluation on 213 incidents."],
    ]
    add_table(doc, "Table 1.1 Research Questions and Evidence", ["ID", "Question", "Evidence Basis"], rq_rows)
    h(doc, "Research Scope and Limitations", 2)
    p(
        doc,
        (
            "The empirical scope is restricted to labeled hijack, leak, and outage "
            "incidents observable through public RouteViews collectors and enriched with "
            "publicly available CAIDA and RPKI data. The system classifies and ranks "
            "incident categories; it does not prove causal mechanisms or replace operator "
            "forensics. Validation and Final Holdout partitions were separated at the "
            "incident level, and the Final Holdout remained sealed until the terminal "
            "evaluation."
        ),
    )
    h(doc, "Praxis Organization", 2)
    for item in [
        "Chapter 1 introduces the motivation, problem, objectives, research questions, scope, and limitations.",
        "Chapter 2 reviews BGP measurement, incident types, RPKI, AS relationships, and machine-learning approaches to routing anomalies.",
        "Chapter 3 describes the incident corpus, RouteViews processing, feature engineering, enrichment, modeling, and evaluation protocol.",
        "Chapter 4 presents the experimental results, feature ablations, model comparisons, validation evidence, K1 architecture, and Final Holdout findings.",
        "Chapter 5 maps findings to the research questions, states contributions and limitations, and identifies future research.",
    ]:
        p(doc, item, style="List Bullet")


def add_chapter2(doc: Document) -> None:
    h(doc, "Chapter 2-Literature Review")
    sections = [
        (
            "Introduction",
            [
                "The literature supporting this Praxis spans BGP operation, routing anomalies, RPKI, AS relationships, public measurement archives, machine learning for network anomalies, and ranking-oriented evaluation. The central theme is that BGP incidents are not only detection problems; once an abnormal routing event is observed, operators benefit from evidence that helps prioritize plausible incident categories.",
                "Prior work has produced important specialized systems for hijack detection, leak detection, anomaly detection, and routing-security measurement. This Praxis builds on that foundation by integrating behavioral, AS-relationship, and RPKI signals into a single incident-level classification and ranking framework.",
            ],
        ),
        (
            "Foundational Studies on BGP Operation and Measurement",
            [
                "BGP-4 defines the interdomain routing protocol used to exchange reachability among autonomous systems (Rekhter et al., 2006). Route selection depends on path attributes, local policy, and incremental UPDATE messages. This design supports Internet-scale policy autonomy, but it also means global behavior emerges from many local decisions rather than from centralized validation.",
                "Public collector projects such as RouteViews and BGPStream provide a practical foundation for historical BGP analysis (University of Oregon, 2000; Orsini et al., 2016). These archives make it possible to reconstruct observed announcements, withdrawals, AS paths, peer visibility, and temporal changes around incidents, while also imposing an important limitation: collector visibility is partial and does not equal enterprise-local telemetry.",
            ],
        ),
        (
            "BGP Routing Incidents: Hijacks, Leaks, and Outages",
            [
                "Prefix hijacks occur when an unauthorized or unexpected AS originates reachability for a prefix. Prior systems such as PHAS, Argus, ARTEMIS, and forged-origin hijack detectors show that hijack detection can use origin changes, prefix relationships, and validation evidence (Lad et al., 2006; Shi et al., 2012; Sermpezis et al., 2018; Holterbach et al., 2024).",
                "Route leaks occur when routes are propagated beyond intended policy boundaries. The literature emphasizes AS-path relationships, role violations, valley-free routing, and leak localization (Azimov et al., 2022; Bagnulo et al., 2022; Li et al., 2023; Xie et al., 2025). These studies motivate CAIDA relationship features and valley-free indicators in this Praxis.",
                "Outages differ from hijacks and leaks because the operational event is often a loss of reachability or service availability rather than a suspicious path assertion. In BGP data, outages may appear through withdrawals, reduced peer visibility, or instability for affected ASNs and prefixes. This distinction matters for both feature design and label interpretation.",
            ],
        ),
        (
            "RPKI, ROA, and Route Origin Validation",
            [
                "RPKI provides a cryptographic framework for expressing which AS is authorized to originate a prefix (Lepinski & Kent, 2012). Route Origin Authorizations and Route Origin Validation can classify observed prefix-origin pairs as valid, invalid, or unknown (Mohapatra et al., 2013; Bush, 2014).",
                "RPKI deployment and operational behavior are uneven across the Internet, so invalidity is not a complete incident explanation by itself (Chung et al., 2019; Rodday et al., 2024). For this Praxis, RPKI is treated as a predictive feature family rather than as a sole label generator. This is especially important for hijack classification, where origin authorization evidence can be highly informative but not universally available.",
            ],
        ),
        (
            "AS Relationships, Routing Policy, and Valley-Free Routing",
            [
                "AS-relationship inference has long been used to characterize routing policy and path economics (Gao, 2001; Gao & Rexford, 2001; Dimitropoulos et al., 2007; Luckie et al., 2013). Customer, provider, and peer relationships provide a policy context for interpreting observed AS paths.",
                "Valley-free assumptions are not perfect, but violations and unusual relationship sequences are useful evidence for route-leak analysis. The CAIDA feature family in this Praxis operationalizes this literature by measuring relationship counts, peer/provider/customer transitions, valley-free violations, and related path-context signals.",
            ],
        ),
        (
            "Machine Learning for BGP Anomaly Detection and Incident Classification",
            [
                "Machine-learning studies have used support vector machines, decision trees, Random Forests, gradient boosting, recurrent neural networks, graph representations, and hybrid robust-statistical methods for BGP anomalies (Al-Rousan & Trajković, 2012; Sanchez et al., 2019; Romo-Chavero et al., 2025). These works show that engineered routing features remain competitive for structured BGP data.",
                "Tree-based methods are particularly appropriate for tabular feature representations. Random Forest provides a strong nonlinear baseline (Breiman, 2001), while XGBoost often performs strongly on tabular classification tasks through sequential boosting and regularization (Chen & Guestrin, 2016). The final system in this Praxis uses XGBoost in both window-level and incident-level stages.",
                "The literature also motivates careful evaluation. Random row-level splitting can leak information when many windows originate from the same incident. This Praxis therefore uses incident-level partitions and grouped cross-validation so that windows from one incident cannot appear in both training and evaluation folds.",
            ],
        ),
        (
            "Explainability, Ranking, and Operational Interpretation",
            [
                "Operational incident support requires more than a single class label. Ranking metrics such as Top-2 accuracy and mean reciprocal rank evaluate whether the true incident category appears near the top of the model's ordered output. These measures fit operator workflows in which the first investigation step may not always be correct, but a short ranked list can still reduce search effort.",
                "Explainability methods such as SHAP can help interpret which features influence predictions (Lundberg & Lee, 2017; Lundberg et al., 2020). However, feature attribution does not prove causality. This Praxis treats explanations as diagnostic aids and evidence checks, not as proof that a model discovered the real-world root cause.",
            ],
        ),
        (
            "Summary of Research Gaps and Study Positioning",
            [
                "The literature leaves several gaps that this Praxis addresses. First, many systems focus on one incident type, while operational triage often requires comparison among multiple plausible incident categories. Second, routing-security signals and AS-relationship signals are often studied separately rather than integrated into one frozen feature corpus. Third, evaluation frequently risks leakage when windows rather than incidents define the independent unit.",
                "This study positions its contribution as an engineering and empirical validation contribution: a reproducible incident-level corpus, a frozen Behavioral+CAIDA+RPKI feature representation, leakage-safe model development, ranked incident-cause output, and a terminal Final Holdout evaluation.",
            ],
        ),
    ]
    for title, paragraphs in sections:
        h(doc, title, 2)
        for paragraph in paragraphs:
            p(doc, paragraph)
    h(doc, "Detailed Literature Review from Current Manuscript", 2)
    p(
        doc,
        (
            "The following section preserves and updates the broader literature-review "
            "material already developed in the current Praxis manuscript. Terminology "
            "has been aligned with the corrected study scope: classification and ranked "
            "incident-cause output rather than unsupported causal root-cause proof."
        ),
    )
    last = None
    for style, text in source_chapter2_paragraphs():
        if text in {
            "Chapter 2—Literature Review",
            "Introduction",
            "Foundational Studies on BGP Operation and Measurement",
            "BGP Routing Incidents and Outages",
            "BGP Security Mechanisms and Routing Validation",
            "AS Relationships, Routing Policies, and Valley-Free Routing",
            "BGP Anomaly Detection Techniques",
            "Machine Learning for Network Fault Detection and Diagnosis",
            "Summary of Research Gaps and Study Positioning",
            "Literature Synthesis and Model Justification",
        }:
            if text != last:
                h(doc, text, 2)
                last = text
            continue
        if text.startswith("Table ") or text.startswith("Figure "):
            continue
        if text != last:
            p(doc, text)
            last = text


def add_chapter3(doc: Document) -> None:
    h(doc, "Chapter 3-Methodology")
    h(doc, "Research Design Overview", 2)
    p(doc, "This Praxis used a quantitative experimental design. Historical BGP incidents were represented as five-minute RouteViews observation windows, enriched with routing-policy and route-origin validation context, and evaluated under incident-level partitions. The independent experimental unit was the incident, not the individual BGP UPDATE message or five-minute feature row.")
    add_caption(doc, "Figure 3.1 Final Research Workflow")
    p(doc, "Labeled incident roster -> incident-level partitioning -> RouteViews MRT acquisition -> five-minute window construction -> Behavioral features -> CAIDA enrichment -> RPKI/ROV enrichment -> V4 feature corpus -> grouped TRAIN development -> Validation confirmation -> TRAIN-only K1 development -> sealed Final Holdout evaluation.")
    h(doc, "Incident Corpus and Partitions", 2)
    p(doc, "The authoritative frozen v2 incident corpus contains 1,419 independent incidents: 473 Hijack, 473 Leak, and 473 Outage incidents. Confirmed duplicate groups were removed before freezing the roster. The final partitions were generated at the incident level and are authoritative for all downstream windows and feature rows.")
    add_table(
        doc,
        "Table 3.1 Incident Partition Summary",
        ["Partition", "Hijack", "Leak", "Outage", "Total", "Role"],
        [
            ["TRAIN", "331", "331", "331", "993", "Development, grouped CV, model selection"],
            ["Validation", "71", "71", "71", "213", "Single predeclared confirmation for Phase J candidate"],
            ["Final Holdout", "71", "71", "71", "213", "One-time terminal evaluation"],
        ],
    )
    p(doc, "The split design follows two methodological rules that recur throughout the pipeline. First, complete incidents are the experimental unit. Second, every downstream five-minute feature row inherits the partition of its parent incident. This prevents a long incident from contributing windows to both training and evaluation and protects the Final Holdout from indirect tuning.")
    p(doc, "The v2 roster remediation process addressed duplicate groups, timestamp semantics, leak-prefix availability, and provenance before freezing the final corpus. Confirmed duplicates were collapsed logically, replacement incidents were selected from larger candidate pools, and the final roster was audited before the v2 split was produced.")
    h(doc, "BGP UPDATE Data Collection and Windowing", 2)
    p(doc, "RouteViews MRT UPDATE archives were used as the public BGP observation source. Incident acquisition windows were derived from incident-specific time bounds and active-window policies. For hijack and leak incidents, detector first-to-last spans could contain silent gaps, so active windows were derived from detector-supported evidence rather than assuming the full span was continuously active. Outage intervals used documented Cloudflare start and end times when they passed quality controls.")
    p(doc, "Raw UPDATE data were parsed into observations, then aggregated into five-minute windows. Five-minute windows were selected because RouteViews archives and prior BGP measurement practice support this cadence, while still providing enough temporal structure for incident-level summarization.")
    p(doc, "A compact enrichment-evidence store was introduced so that repeated prefix, origin, and AS-path observations could be aggregated without retaining the full raw parsed corpus. Each evidence row represents a unique window, prefix, origin ASN, normalized AS path, and collector combination with an observation count preserving multiplicity. This design allowed CAIDA and RPKI/ROV enrichment to be recalculated without reparsing all raw MRT files.")
    h(doc, "Feature Engineering", 2)
    p(doc, "The final V4 feature representation contains 74 features. Behavioral features quantify UPDATE volume, withdrawals, announcements, AS-path properties, prefix visibility, origin changes, and baseline deviations. CAIDA features add AS-relationship and valley-free context. RPKI/ROV features add historical route-origin authorization and validation status.")
    add_table(
        doc,
        "Table 3.2 Frozen Feature Families",
        ["Feature family", "Count", "Purpose"],
        [
            ["Behavioral BGP", "44", "Temporal routing behavior from five-minute RouteViews windows"],
            ["CAIDA AS relationship", "14", "Provider/peer/customer and valley-free path context"],
            ["RPKI/ROV", "16", "Route-origin authorization and validation context"],
            ["Total V4", "74", "Final frozen model-ready feature representation"],
        ],
    )
    add_table(
        doc,
        "Table 3.3 Feature Variant Definitions",
        ["Variant", "Feature count", "Definition"],
        [
            ["V1", "44", "Behavioral features only"],
            ["V2", "58", "Behavioral plus CAIDA features"],
            ["V3", "60", "Behavioral plus selected RPKI features"],
            ["V4", "74", "Behavioral plus CAIDA plus full RPKI/ROV features"],
        ],
    )
    p(doc, "Behavioral features were frozen first so that later enrichment did not change the original routing-behavior representation. CAIDA enrichment added relationship-aware path context, while RPKI/ROV enrichment mapped each prefix-origin observation to the historical authorization state appropriate for the observation time. The V4 schema was then frozen at exactly 74 features.")
    p(doc, "Feature extraction was intentionally separated from modeling. No Validation or Final Holdout information was used to choose the V4 feature schema, imbalance treatment, K1 architecture, or final model hyperparameters.")
    h(doc, "Normal Label Methodology", 2)
    p(doc, "Normal-window construction was treated conservatively. Only windows that satisfied rule-based stability criteria and did not overlap verified incident windows were considered candidate Normal examples. The Normal-vs-Incident Stage 1 experiments remained unresolved and auxiliary; the final K1 architecture does not depend on a production Stage 1 threshold.")
    h(doc, "Model Development Protocol", 2)
    p(doc, "TRAIN data supported iterative model development through grouped cross-validation. Grouping by incident prevented windows from the same incident from crossing folds. Validation was opened once in Phase J for a predeclared refined XGBoost candidate. K1 was subsequently developed as a versioned TRAIN-only post-Validation experiment and was not evaluated on Validation.")
    add_table(
        doc,
        "Table 3.4 Modeling Phases",
        ["Phase", "Purpose", "Partition access"],
        [
            ["Baseline and imbalance phases", "Establish RF/XGB baselines and I0/I1/I2 weighting effects", "TRAIN only"],
            ["Feature ablation", "Evaluate V1-V4 and feature-family contribution", "TRAIN only"],
            ["Phase J", "Single predeclared Validation confirmation of refined XGBoost", "TRAIN + Validation"],
            ["K1/K2/K2.5/L0", "TRAIN-only incident-level stacking, stability checks, architecture comparison, final freeze", "TRAIN only"],
            ["Phase L", "One-time terminal Final Holdout evaluation", "Final Holdout only after freeze"],
        ],
    )
    h(doc, "Final K1 Architecture", 2)
    p(doc, "The terminal model is the A0_XGB_TO_XGB K1 architecture. Level 1 uses XGBoost window-level probability models. Those probability outputs are summarized at the incident level using temporal and probability-distribution meta-features. A second-level XGBoost model then classifies and ranks the incident category. The frozen second-level model hash is 40df040321ef.")
    add_caption(doc, "Figure 3.2 Final K1 Incident-Level Stacking Architecture")
    p(doc, "Five-minute V4 feature windows -> Level-1 XGBoost probability outputs -> incident-level temporal/probability meta-features -> Level-2 XGBoost classifier -> ranked Hijack/Leak/Outage output.")
    h(doc, "Evaluation Metrics", 2)
    p(doc, "The primary incident-level metrics were Macro F1, Top-1 accuracy, Top-2 accuracy, and mean reciprocal rank. Per-class precision, recall, and F1 were reported for Hijack, Leak, and Outage. Final Holdout uncertainty was estimated at the incident level; window-level uncertainty was not mixed with incident-level uncertainty.")
    h(doc, "Evaluation Integrity and Holdout Protection", 2)
    p(doc, "The Final Holdout was sealed during feature selection, model selection, threshold experiments, architecture comparison, balancing decisions, and hyperparameter tuning. No post-Holdout model changes were made. This protects the terminal result as the study's independent generalization estimate.")
    add_table(
        doc,
        "Table 3.5 Evaluation-Integrity Controls",
        ["Risk", "Control"],
        [
            ["Incident/window leakage", "Incident-level partitions and grouped cross-validation"],
            ["Validation overuse", "Single predeclared Phase J confirmation"],
            ["Post-Validation model development", "Versioned TRAIN-only K1/K2/K2.5/L0 record"],
            ["Holdout tuning", "Final Holdout remained sealed until Phase L"],
            ["Window-count dominance", "Incident-aware weighting and incident-level metrics"],
        ],
    )
    h(doc, "Methodological Limitations", 2)
    p(doc, "The methodology relies on public RouteViews collectors, so visibility may differ from enterprise-local routing state. Incident labels originate from external sources and inherit their documentation limits. Leak events do not always include exact affected prefix lists in the source data, but they remain identifiable by time, leak ASN, AS-path segment, and origin/path evidence. Stage 1 Normal-vs-Incident discrimination remained auxiliary.")


def add_chapter4(doc: Document) -> None:
    h(doc, "Chapter 4-Results")
    h(doc, "Overview of Experimental Results", 2)
    p(doc, "Chapter 4 reports the evidence trail from baseline modeling through the terminal Final Holdout evaluation. The results should be interpreted in chronological order: TRAIN grouped-CV development first, one-time Validation confirmation second, post-Validation TRAIN-only K1 development third, and sealed Final Holdout evaluation last.")
    add_table(
        doc,
        "Table 4.1 Model Performance Progression",
        ["Stage", "Evaluation design", "Feature set", "Incident Macro F1", "Status"],
        [
            ["RF baseline", "TRAIN grouped CV", "Behavioral V1", "0.463174", "Historical baseline"],
            ["Refined RF H2", "TRAIN grouped CV", "V4/I1", "0.784588", "RF benchmark"],
            ["Refined XGBoost I2", "TRAIN grouped CV", "V4/I1", "0.805678", "Promoted to Validation"],
            ["Phase J Validation", "Validation confirmation", "V4/I1", "0.713904", "Passed floors"],
            ["K1 stacker", "TRAIN nested CV", "V4/I1", "0.934666", "Post-Validation TRAIN development"],
            ["K2 multi-seed K1", "TRAIN nested CV", "V4/I1", "0.934304 ± 0.002309", "Forensic/stability"],
            ["L0 fixed deployable K1", "TRAIN nested CV", "V4/I1", "0.931653 ± 0.002150", "Pre-Holdout frozen"],
            ["Final Holdout", "One-time terminal", "V4/I1/K1", "0.920177", "Terminal"],
        ],
    )
    h(doc, "Feature Variant and Feature-Family Evidence", 2)
    p(doc, "Feature-variant and feature-family experiments supported the value of adding routing-policy and origin-validation context to behavioral BGP features. RPKI/ROV features contributed strongly to Hijack classification, while CAIDA and peer-visibility context provided useful signal for Leak and Outage separation. These are predictive feature-contribution findings, not causal claims about the underlying incidents.")
    h(doc, "Random Forest and XGBoost Model Selection", 2)
    p(doc, "Random Forest remained an important benchmark, but refined XGBoost outperformed the refined Random Forest comparator under TRAIN grouped-CV. The refined Random Forest reference produced incident Macro F1 of 0.784588, while the refined XGBoost candidate produced incident Macro F1 of 0.805678.")
    add_table(
        doc,
        "Table 4.2 Refined RF and Refined XGBoost TRAIN Grouped-CV Comparison",
        ["Model", "Incident Macro F1", "Top-1", "Top-2", "MRR", "Disposition"],
        [
            ["Refined Random Forest", "0.784588", "0.787513", "0.932931", "0.882578", "Benchmark"],
            ["Refined XGBoost", "0.805678", "0.807049", "0.934743", "0.892649", "Promoted to Phase J"],
        ],
    )
    h(doc, "Phase J Validation Confirmation", 2)
    p(doc, "Phase J was the first authorized access to the Validation partition. The predeclared refined XGBoost candidate passed the frozen Validation floors, but the later K1 architecture was not evaluated on Validation. Therefore, Validation evidence confirms the pre-K1 refined XGBoost candidate and does not make K1 a Validation-confirmed model.")
    add_table(
        doc,
        "Table 4.3 Phase J Validation Metrics",
        ["Metric", "Validation value"],
        [
            ["Incident Macro F1", "0.713904"],
            ["Top-1", "0.718310"],
            ["Top-2", "0.901408"],
            ["MRR", "0.842723"],
            ["Hijack F1", "0.634146"],
            ["Leak F1", "0.774648"],
            ["Outage F1", "0.732919"],
        ],
    )
    h(doc, "K1 Incident-Level Stacking Results", 2)
    p(doc, "The K1 architecture materially improved incident-level classification and ranking compared with mean probability aggregation by allowing the second-level model to learn temporal and probability-distribution summaries across windows. K2 multi-seed stability and K2.5 architecture comparisons supported keeping the XGBoost-to-XGBoost K1 architecture.")
    add_table(
        doc,
        "Table 4.4 K1 Stability and Final Pre-Holdout Reference",
        ["Experiment", "Incident Macro F1", "Std", "Interpretation"],
        [
            ["K2 multi-seed K1 stability", "0.934304", "0.002309", "TRAIN-only stability evidence"],
            ["Fixed deployable K1", "0.931653", "0.002150", "Frozen pre-Holdout reference"],
        ],
    )
    h(doc, "Final Holdout Performance", 2)
    p(doc, "The Final Holdout was accessed once after the K1 deployable model was frozen. It contained 213 incidents, balanced across 71 Hijack, 71 Leak, and 71 Outage incidents, and 14,906 V4 window rows. The final result is the primary independent generalization estimate for this Praxis.")
    add_table(
        doc,
        "Table 4.5 Final Holdout Performance",
        ["Metric", "Value", "Sample unit"],
        [
            ["Incident Macro F1", "0.920177", "Incident"],
            ["95% CI for Macro F1", "[0.880117, 0.954402]", "Incident bootstrap"],
            ["Top-1 Accuracy", "0.920188", "Incident"],
            ["Top-2 Accuracy", "0.990610", "Incident"],
            ["Mean Reciprocal Rank", "0.958529", "Incident"],
        ],
    )
    add_table(
        doc,
        "Table 4.6 Final Holdout Per-Class F1",
        ["Class", "F1"],
        [["Hijack", "0.877698"], ["Leak", "0.904110"], ["Outage", "0.978723"]],
    )
    add_table(
        doc,
        "Table 4.7 Final Holdout Ranking Distribution",
        ["True-class rank", "Incident count"],
        [["Rank 1", "196"], ["Rank 2", "15"], ["Rank 3", "2"], ["Top-2 total", "211 of 213"]],
    )
    h(doc, "Class-Specific Interpretation", 2)
    p(doc, "Outage was the strongest final class, with F1 of 0.978723. Leak achieved F1 of 0.904110, indicating strong separation after Behavioral, CAIDA, and RPKI/ROV enrichment. Hijack remained the most difficult final class, with F1 of 0.877698, consistent with the complexity of distinguishing unauthorized origin behavior, sparse visibility, and overlapping path changes.")
    h(doc, "Runtime and Operational Feasibility", 2)
    p(doc, "The measured warm K1 downstream classification path was approximately 25.58 milliseconds per incident after features and Level-1 probabilities were available. This measurement should not be described as end-to-end live detection time because it excludes live BGP stream ingestion, parsing, feature construction, and production integration. It supports the narrower claim that the trained downstream ranking step is computationally lightweight.")
    h(doc, "Negative and Null Findings", 2)
    p(doc, "Several negative results strengthened the final design. Stage 1 Normal-vs-Incident discrimination did not become the final architecture because Normal precision/recall tradeoffs remained weak. Domain-expert model variants were not adopted. Class balancing alone was not sufficient. Threshold tuning was not sufficient. These results reduced the candidate space and made the final K1 selection more defensible.")
    h(doc, "Chapter Summary", 2)
    p(doc, "The terminal evidence shows strong incident-level generalization on the untouched Final Holdout, with high ranking quality and especially strong outage performance. The strongest claim supported by the evidence is predictive classification and ranked incident-cause output from public BGP routing telemetry, not causal proof or enterprise-local root-cause diagnosis.")


def add_chapter5(doc: Document) -> None:
    h(doc, "Chapter 5-Conclusions")
    h(doc, "Overview", 2)
    p(doc, "This Praxis developed and evaluated a reproducible machine-learning pipeline for classifying and ranking BGP hijack, leak, and outage incidents. The final system combined behavioral BGP features, CAIDA AS-relationship context, and RPKI/ROV route-origin validation features in an incident-level XGBoost stacking architecture.")
    h(doc, "Mapping of Findings to Research Questions and Hypotheses", 2)
    add_table(
        doc,
        "Table 5.1 Research Question Findings",
        ["Research question", "Finding", "Disposition"],
        [
            ["RQ1", "Historical RouteViews data supported a reproducible incident corpus and V4 feature representation.", "Supported"],
            ["RQ2", "CAIDA and RPKI/ROV enrichments provided useful predictive signal, with RPKI especially relevant to Hijack.", "Supported with caveat"],
            ["RQ3", "K1 XGBoost-to-XGBoost stacking was selected after TRAIN-only architecture comparison and forensic checks.", "Supported"],
            ["RQ4", "The locked K1 model generalized strongly on the untouched Final Holdout.", "Supported"],
        ],
    )
    h(doc, "Key Contributions", 2)
    for item in [
        "A reproducible BGP incident feature-engineering pipeline using historical RouteViews MRT data.",
        "A frozen 74-feature Behavioral+CAIDA+RPKI/ROV representation for hijack, leak, and outage classification.",
        "Incident-aware partitioning, grouping, and weighting controls that address unequal window counts and leakage risk.",
        "A leakage-safe incident-level temporal stacking architecture that produces both class predictions and ranked incident-cause output.",
        "A terminal one-time Final Holdout evaluation on 213 independent incidents.",
    ]:
        p(doc, item, style="List Bullet")
    h(doc, "Limitations", 2)
    for item in [
        "Public RouteViews collectors provide partial interdomain visibility and may not reflect enterprise-local routing state.",
        "Incident labels come from external incident sources and inherit their documentation quality limits.",
        "Normal-vs-Incident Stage 1 discrimination remained unresolved and auxiliary.",
        "K1 was developed after Validation, so the Final Holdout is the primary independent generalization estimate for the terminal K1 model.",
        "The Final Holdout contains 213 incidents; confidence intervals should be interpreted with that sample size in mind.",
        "The measured 25.58 ms downstream K1 runtime excludes live ingestion, parsing, and feature computation.",
    ]:
        p(doc, item, style="List Bullet")
    h(doc, "Recommendations for Future Research", 2)
    for item in [
        "Develop a stronger real-time Normal-vs-Incident formulation using streaming incremental baselines and operator-validated benign periods.",
        "Evaluate shorter observation horizons such as 30 seconds, 1 minute, 2 minutes, and 3 minutes after rebuilding compatible features from higher-resolution evidence.",
        "Integrate enterprise-local telemetry, including router logs, BMP feeds, SNMP, flow telemetry, and operator ticket outcomes, to connect public BGP symptoms with local operational context.",
        "Extend the corpus with additional verified incidents, collectors, and geographic perspectives.",
        "Study calibrated ranked probabilities and operator diagnostic-time impact through controlled user studies.",
        "Explore graph-based AS representations and sequence models once larger leakage-safe incident corpora are available.",
    ]:
        p(doc, item, style="List Bullet")
    h(doc, "Conclusion", 2)
    p(doc, "The final evidence supports the central claim that BGP incident classification and ranked incident-cause output can be learned from a carefully controlled combination of RouteViews behavior, AS-relationship context, and RPKI/ROV signals. The terminal K1 system achieved Macro F1 of 0.920177 and ranked the true class in the Top-2 for 211 of 213 Final Holdout incidents. These findings support the Praxis as an engineering contribution to BGP incident analysis while preserving clear boundaries around causality, public-collector visibility, and unresolved Normal detection.")


def add_references(doc: Document) -> None:
    h(doc, "References")
    refs = current_references()
    seen = set()
    for ref in refs:
        normalized = re.sub(r"\s+", " ", ref).strip()
        if not normalized or normalized in seen:
            continue
        seen.add(normalized)
        para = p(doc, normalized)
        para.paragraph_format.first_line_indent = Inches(-0.5)
        para.paragraph_format.left_indent = Inches(0.5)


def normalize_styles(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.25)
    try:
        normal = doc.styles["Normal"]
        normal.font.name = "Times New Roman"
        normal.font.size = Pt(12)
    except Exception:
        pass


def build_docx() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document(TEMPLATE)
    clear_document(doc)
    normalize_styles(doc)
    add_front_matter(doc)
    add_chapter1(doc)
    add_page_break(doc)
    add_chapter2(doc)
    add_page_break(doc)
    add_chapter3(doc)
    add_page_break(doc)
    add_chapter4(doc)
    add_page_break(doc)
    add_chapter5(doc)
    add_page_break(doc)
    add_references(doc)
    doc.save(DOCX_OUT)


def export_pdf() -> None:
    if not SOFFICE.exists():
        raise FileNotFoundError(f"LibreOffice executable not found: {SOFFICE}")
    profile_dir = Path(tempfile.mkdtemp(prefix="lo-praxis-profile-"))
    subprocess.run(
        [
            str(SOFFICE),
            "--headless",
            f"-env:UserInstallation=file://{profile_dir}",
            "--convert-to",
            "pdf",
            "--outdir",
            str(OUT_DIR),
            str(DOCX_OUT),
        ],
        check=True,
        cwd=str(OUT_DIR),
    )
    generated = OUT_DIR / f"{DOCX_OUT.stem}.pdf"
    if generated != PDF_OUT and generated.exists():
        shutil.move(generated, PDF_OUT)


def main() -> None:
    build_docx()
    export_pdf()
    print(DOCX_OUT)
    print(PDF_OUT)


if __name__ == "__main__":
    main()
