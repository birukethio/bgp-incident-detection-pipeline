from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path("/Users/biruk/Documents/New project")
OUT_DIR = ROOT / "deliverables/chapter4_tables"
OUT_DOCX = OUT_DIR / "Table_4_9_RF_XGBoost_Performance_Comparison.docx"


ROWS = [
    [
        "Phase H2\nRefined RF",
        "0.784588",
        "0.001342",
        "0.787513",
        "0.932931",
        "0.882578",
        "0.721923",
        "0.841110",
        "0.790732",
        "Fit: 399.192 s\nInfer: 1.321 s",
    ],
    [
        "Phase I\nXGBoost",
        "0.799599",
        "0.000000",
        "0.800604",
        "0.933535",
        "0.889225",
        "0.751269",
        "0.848580",
        "0.798949",
        "Fit: 21.699 s\nInfer: 0.158 s",
    ],
    [
        "Phase I2\nRefined XGBoost",
        "0.805678",
        "0.002042",
        "0.807049",
        "0.934743",
        "0.892649",
        "0.756196",
        "0.851923",
        "0.808913",
        "Fit: 32.989 s\nInfer: 0.215 s",
    ],
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_inches: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def style_paragraph(paragraph, size=8, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER):
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(size)
        run.bold = bold


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    title = doc.add_paragraph("Table 4.9 Refined Random Forest and XGBoost Performance Comparison")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    for run in title.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.bold = True

    headers = [
        "Model\nStage",
        "Macro F1",
        "SD",
        "Top-1",
        "Top-2",
        "MRR",
        "Hijack F1",
        "Leak F1",
        "Outage F1",
        "Runtime\nFit / Inference",
    ]
    widths = [1.05, 0.78, 0.75, 0.72, 0.72, 0.72, 0.76, 0.76, 0.78, 1.25]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False

    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, "D9EAF7")
        set_cell_margins(cell)
        set_cell_width(cell, widths[i])
        for p in cell.paragraphs:
            style_paragraph(p, size=7.5, bold=True)

    for row in ROWS:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cells[i])
            set_cell_width(cells[i], widths[i])
            for p in cells[i].paragraphs:
                style_paragraph(p, size=7.5, bold=False)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(8)
    note.paragraph_format.space_after = Pt(0)
    note.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = note.add_run(
        "Note. Results are TRAIN-only grouped-CV values under the frozen V4/I1 framework. "
        "Runtime is reported as fit seconds and inference seconds. Source: docs/chapter4_xgboost_comparison_results_v1.csv."
    )
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)

    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    main()
