from __future__ import annotations

import csv
import re
import sys
import zipfile
from collections import Counter, defaultdict
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document


W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def text_of_para_xml(p: ET.Element) -> str:
    parts = []
    for t in p.iter(f"{W_NS}t"):
        if t.text:
            parts.append(t.text)
    return "".join(parts)


def write_csv(path: Path, rows: list[dict], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def extract_doc_xml(path: Path) -> ET.Element:
    with zipfile.ZipFile(path) as zf:
        return ET.fromstring(zf.read("word/document.xml"))


def para_page_breaks(path: Path) -> list[dict]:
    root = extract_doc_xml(path)
    rows = []
    page = 1
    idx = 0
    for p in root.iter(f"{W_NS}p"):
        idx += 1
        text = text_of_para_xml(p).strip()
        style = ""
        ppr = p.find(f"{W_NS}pPr")
        if ppr is not None:
            ps = ppr.find(f"{W_NS}pStyle")
            if ps is not None:
                style = ps.attrib.get(f"{W_NS}val", "")
        if text:
            rows.append({"paragraph_index": idx, "page_estimate": page, "style": style, "text": text[:500]})
        for br in p.iter(f"{W_NS}br"):
            if br.attrib.get(f"{W_NS}type") == "page":
                page += 1
        for _last in p.iter(f"{W_NS}lastRenderedPageBreak"):
            page += 1
    return rows


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("usage: praxis_docx_audit.py input.docx output_dir")
    docx_path = Path(sys.argv[1])
    out_dir = Path(sys.argv[2])
    out_dir.mkdir(parents=True, exist_ok=True)
    doc = Document(docx_path)

    para_rows = []
    heading_rows = []
    caption_rows = []
    for i, p in enumerate(doc.paragraphs, start=1):
        text = p.text.strip()
        style = p.style.name if p.style is not None else ""
        if not text:
            continue
        para_rows.append(
            {
                "paragraph_index": i,
                "style": style,
                "word_count": len(re.findall(r"[A-Za-z0-9]+", text)),
                "text": text[:1000],
            }
        )
        if style.startswith("Heading") or re.match(r"^Chapter\s+\d+", text, re.I):
            heading_rows.append({"paragraph_index": i, "style": style, "text": text})
        if re.match(r"^(Table|Figure)\s+\d+", text, re.I):
            caption_rows.append({"paragraph_index": i, "style": style, "caption_type": text.split()[0], "text": text})

    write_csv(out_dir / "paragraphs.csv", para_rows, ["paragraph_index", "style", "word_count", "text"])
    write_csv(out_dir / "headings.csv", heading_rows, ["paragraph_index", "style", "text"])
    write_csv(out_dir / "captions.csv", caption_rows, ["paragraph_index", "style", "caption_type", "text"])
    write_csv(out_dir / "page_estimate_paragraphs.csv", para_page_breaks(docx_path), ["paragraph_index", "page_estimate", "style", "text"])

    full_text = "\n".join(p.text for p in doc.paragraphs)
    citation_re = re.compile(r"\(([A-Z][A-Za-zÀ-ÿ'’\-]+(?:\s+et\s+al\.)?(?:,\s*[A-Z][A-Za-zÀ-ÿ'’\-]+(?:\s+et\s+al\.)?)*,\s*(?:19|20)\d{2}[a-z]?(?:;\s*[^()]+?)*)\)")
    citations = sorted(set(m.group(1) for m in citation_re.finditer(full_text)))
    write_csv(out_dir / "in_text_citations_raw.csv", [{"citation": c} for c in citations], ["citation"])

    ref_start = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().lower() in {"references", "reference"}:
            ref_start = i
            break
    ref_rows = []
    if ref_start is not None:
        for i, p in enumerate(doc.paragraphs[ref_start + 1 :], start=ref_start + 2):
            text = p.text.strip()
            if text:
                ref_rows.append({"paragraph_index": i, "text": text})
    write_csv(out_dir / "references_raw.csv", ref_rows, ["paragraph_index", "text"])

    styles = Counter(row["style"] for row in para_rows)
    write_csv(out_dir / "style_counts.csv", [{"style": k, "count": v} for k, v in styles.most_common()], ["style", "count"])

    chapter_counts = defaultdict(int)
    current = "front_matter"
    for row in para_rows:
        text = row["text"]
        m = re.match(r"^(Chapter\s+\d+)", text, re.I)
        if m:
            current = m.group(1).title()
        chapter_counts[current] += int(row["word_count"])
    write_csv(out_dir / "word_counts_by_chapter_estimate.csv", [{"chapter": k, "word_count": v} for k, v in chapter_counts.items()], ["chapter", "word_count"])

    sections = doc.sections
    section_rows = []
    for i, section in enumerate(sections, start=1):
        section_rows.append(
            {
                "section": i,
                "page_width_inches": f"{section.page_width.inches:.3f}",
                "page_height_inches": f"{section.page_height.inches:.3f}",
                "top_margin_inches": f"{section.top_margin.inches:.3f}",
                "bottom_margin_inches": f"{section.bottom_margin.inches:.3f}",
                "left_margin_inches": f"{section.left_margin.inches:.3f}",
                "right_margin_inches": f"{section.right_margin.inches:.3f}",
            }
        )
    write_csv(out_dir / "section_layout.csv", section_rows, ["section", "page_width_inches", "page_height_inches", "top_margin_inches", "bottom_margin_inches", "left_margin_inches", "right_margin_inches"])


if __name__ == "__main__":
    main()
